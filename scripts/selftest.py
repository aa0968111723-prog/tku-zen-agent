"""自我檢查：不需要 API 金鑰，確認知識庫檢索與四個產出工具都正常。

    python scripts/selftest.py
"""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from app import config, retrieval, tools  # noqa: E402

PASS, FAIL = "  [OK] ", "  [!!] "
failures = 0


def check(label: str, ok: bool, note: str = "") -> None:
    global failures
    print((PASS if ok else FAIL) + label + (f" — {note}" if note else ""))
    if not ok:
        failures += 1


print("=" * 62)
print("  淡江大學領袖禪學社 · AI 代理 —— 自我檢查")
print("=" * 62)

# 1. 設定
print("\n[1] 設定")
for problem in config.missing_config():
    print(f"  [--] {problem}")
print(f"  產出資料夾：{config.OUTPUT_DIR}")
print(f"  預設落點：{config.DEFAULT_DESTINATION}")

# 2. 知識庫
print("\n[2] 知識庫檢索")
idx = retrieval.get_index(rebuild=True)
stats = idx.stats()
check(f"索引建立：{stats['段落數']} 段", stats["段落數"] > 20)
print(f"       來源：{', '.join(stats['來源'])}")

for query, expect in [
    ("招生管道有哪些", "路宣"),
    ("青年領袖十二項特質", "體魄"),
    ("社課籌備 SOP 一籌二籌", "二籌"),
    ("期初茶會", "茶會"),
    ("經費預算要怎麼分類", "核銷"),
    ("對外文宣不能寫什麼", "療效"),
]:
    hits = idx.search(query, k=5)
    found = any(expect in c.text for _, c in hits)
    check(f'查「{query}」', found, "" if found else f"前 5 名裡找不到「{expect}」")

# 有匯入歷年檔案時，順便確認精選劇本不會被 500 多份檔案淹掉
if any(c.tier >= retrieval.TIER_CURATED for c in idx.chunks) and any(
    c.tier < retrieval.TIER_CURATED for c in idx.chunks
):
    worst = []
    for query in ["社課當天流程時間表", "招生要怎麼規劃", "營隊安全", "評鑑要交什麼"]:
        hits = idx.search(query, k=5)
        n = sum(1 for _, c in hits if c.tier >= retrieval.TIER_CURATED)
        if n < 1:
            worst.append(query)
    check("歷年檔案不會把精選劇本擠出結果", not worst, f"這些查詢完全沒有劇本：{worst}")

# 3. 產出工具
print("\n[3] 產出工具")

r = tools.dispatch("search_knowledge", {"query": "招生管道", "top_k": 3})
check("search_knowledge", r.get("ok") and "路宣" in r.get("message", ""), r.get("message", "")[:80])

r = tools.dispatch("search_knowledge", {"query": "紅燒牛肉麵的湯頭怎麼熬", "top_k": 3})
check(
    "完全查不到時會明確要求不要編造",
    r.get("ok") and r.get("low_confidence") and "不要自行編造" in r.get("message", ""),
    r.get("message", "")[:90],
)

r = tools.dispatch(
    "create_spreadsheet",
    {
        "filename": "_自我檢查_試算表",
        "summary": "自我檢查用",
        "sheets": [
            {
                "name": "社課排程",
                "csv": "週次,日期,主題,帶課者,人數\n第1週,待填,期初茶會,待邀,35\n第2週,待填,改變從認識自己開始,待邀,28",
            },
            {"name": "預算", "csv": "項目,單價,數量,小計\n海報,120,10,=B2*C2\n總計,,,=SUM(D2:D2)"},
        ],
    },
)
check("create_spreadsheet", r.get("ok"), r.get("message", "")[:120])
if r.get("local_path"):
    p = Path(r["local_path"])
    check("  檔案確實存在且不是空的", p.exists() and p.stat().st_size > 3000, f"{p.stat().st_size if p.exists() else 0} bytes")
    try:
        from openpyxl import load_workbook

        wb = load_workbook(p)
        check("  可以正常開啟", wb.sheetnames == ["社課排程", "預算"], str(wb.sheetnames))
        check("  數字有存成數值型別", isinstance(wb["社課排程"]["E2"].value, int), repr(wb["社課排程"]["E2"].value))
        check("  公式有保留", str(wb["預算"]["D2"].value).startswith("="), repr(wb["預算"]["D2"].value))
    except Exception as exc:  # noqa: BLE001
        check("  可以正常開啟", False, str(exc))

r = tools.dispatch(
    "create_document",
    {
        "filename": "_自我檢查_文件",
        "title": "自我檢查文件",
        "markdown": "# 第一章\n\n這是**粗體**測試。\n\n- 項目一\n- 項目二\n\n| 欄A | 欄B |\n|---|---|\n| 1 | 2 |\n\n> 引言測試\n",
        "also_markdown": True,
    },
)
check("create_document", r.get("ok"), r.get("message", "")[:120])
if r.get("local_path"):
    p = Path(r["local_path"])
    check("  檔案確實存在", p.exists() and p.stat().st_size > 5000, f"{p.stat().st_size if p.exists() else 0} bytes")
    try:
        from docx import Document

        d = Document(p)
        check("  可以正常開啟", len(d.paragraphs) > 3 and len(d.tables) == 1)
    except Exception as exc:  # noqa: BLE001
        check("  可以正常開啟", False, str(exc))

r = tools.dispatch(
    "create_slides",
    {
        "filename": "_自我檢查_簡報",
        "title": "自我檢查",
        "subtitle": "淡江大學領袖禪學社",
        "slides_markdown": "## 第一頁\n- 重點一\n- 重點二\n  - 次層\n\n## 第二頁\n- 另一個重點",
    },
)
check("create_slides", r.get("ok"), r.get("message", "")[:120])
if r.get("local_path"):
    p = Path(r["local_path"])
    try:
        from pptx import Presentation

        prs = Presentation(p)
        check("  可以正常開啟", len(prs.slides) == 3, f"{len(prs.slides)} 頁")
    except Exception as exc:  # noqa: BLE001
        check("  可以正常開啟", False, str(exc))

r = tools.dispatch(
    "create_google_form",
    {
        "filename": "_自我檢查_表單",
        "form_title": "自我檢查表單",
        "description": "測試用",
        "questions": [
            {"type": "short_text", "title": "姓名", "required": True},
            {"type": "multiple_choice", "title": "怎麼知道我們的", "options": ["路宣", "朋友介紹", "IG"]},
            {"type": "scale", "title": "整體評分", "scale_min": 1, "scale_max": 5},
            {"type": "paragraph", "title": "還想說什麼"},
        ],
    },
)
check("create_google_form", r.get("ok"), r.get("message", "")[:120])
if r.get("local_path"):
    gs = Path(r["local_path"]).read_text(encoding="utf-8")
    check("  腳本含建立表單指令", "FormApp.create" in gs and "addScaleItem" in gs)
    check("  中文有正確寫入", "姓名" in gs and "路宣" in gs)

# 4. 容錯：模型亂傳參數
print("\n[4] 容錯（模擬模型傳錯格式）")
r = tools.dispatch("create_spreadsheet", {"filename": "_自我檢查_容錯", "sheets": "欄1,欄2\n值1,值2"})
check("sheets 傳成字串也能救回來", r.get("ok"), r.get("message", "")[:80])

r = tools.dispatch("create_spreadsheet", {"filename": "_自我檢查_容錯2"})
check("缺少必要參數會給出可讀的錯誤", not r.get("ok") and "sheets" in r.get("message", ""))

r = tools.dispatch("不存在的工具", {})
check("呼叫不存在的工具不會炸掉", not r.get("ok"))

# 5. 完整代理迴圈（用假的模型回應，不需要 API 金鑰）
print("\n[5] 代理迴圈（模擬模型，不需金鑰）")


def _agent_loop_test() -> tuple[bool, str]:
    import asyncio

    from app import agent as agent_mod
    from app.llm import Reply, ToolCall

    scripted = [
        # 第一輪：查知識庫
        Reply(tool_calls=[ToolCall(id="c1", name="search_knowledge", arguments={"query": "招生管道"})]),
        # 第二輪：把結果做成試算表（arguments 故意用字串，模擬開源模型的常見行為）
        Reply(
            tool_calls=[
                ToolCall(
                    id="c2",
                    name="create_spreadsheet",
                    arguments={
                        "filename": "_自我檢查_迴圈",
                        "sheets": '[{"name":"招生管道","csv":"管道,負責組別\\n路宣,招生組\\n個接,全體"}]',
                    },
                )
            ]
        ),
        # 第三輪：收尾
        Reply(content="做好了，檔案是 _自我檢查_迴圈.xlsx。"),
    ]
    calls = {"n": 0}

    async def fake_chat(self, messages, tools=None, **kw):  # noqa: ANN001, ARG001
        i = calls["n"]
        calls["n"] += 1
        return scripted[min(i, len(scripted) - 1)]

    original = agent_mod.NvidiaClient.chat
    agent_mod.NvidiaClient.chat = fake_chat  # type: ignore[method-assign]
    try:
        events = []

        async def drive():
            async for ev in agent_mod.run_turn("selftest", "幫我做招生管道表", destination="local"):
                events.append(ev)

        asyncio.run(drive())
    finally:
        agent_mod.NvidiaClient.chat = original  # type: ignore[method-assign]
        agent_mod.reset_session("selftest")

    kinds = [e["type"] for e in events]
    artifacts = [e for e in events if e.get("artifact")]
    finals = [e for e in events if e["type"] == "message"]
    ok = (
        kinds.count("tool_start") == 2
        and kinds.count("tool_end") == 2
        and len(artifacts) == 1
        and len(finals) == 1
        and "done" in kinds
        and not any(e["type"] == "error" for e in events)
    )
    return ok, f"事件序列 {kinds}"


try:
    ok, note = _agent_loop_test()
    check("查知識庫 → 產出檔案 → 收尾，完整跑通", ok, note)
except Exception as exc:  # noqa: BLE001
    check("查知識庫 → 產出檔案 → 收尾，完整跑通", False, f"{type(exc).__name__}: {exc}")

print("\n" + "=" * 62)
if failures:
    print(f"  {failures} 項檢查沒過。")
else:
    print("  全部通過。")
print("=" * 62)
sys.exit(1 if failures else 0)
