"""產出驗證與自動修正。

原則：**驗證不靠模型自覺**。把 verify 做成 orchestrator 的固定階段，
而不是暴露一個 verify_artifact 工具期待模型自己想到要呼叫 ——
開源模型多半不會呼叫，那等於沒有驗證。

檢查分兩種嚴重度：
  · error   —— 一定要修（編造今年事實、療效宣稱、缺必要章節）
  · warning —— 提醒但不擋（可以更好，但不算錯）
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from ..services import current_term as term_service


@dataclass
class Issue:
    rule: str
    severity: str          # error | warning
    message: str
    fix_hint: str = ""

    def to_dict(self) -> dict[str, str]:
        return {"rule": self.rule, "severity": self.severity, "message": self.message, "fix_hint": self.fix_hint}


@dataclass
class Report:
    filename: str
    kind: str
    issues: list[Issue] = field(default_factory=list)
    checked: list[str] = field(default_factory=list)
    text_length: int = 0

    @property
    def errors(self) -> list[Issue]:
        return [i for i in self.issues if i.severity == "error"]

    @property
    def warnings(self) -> list[Issue]:
        return [i for i in self.issues if i.severity == "warning"]

    @property
    def ok(self) -> bool:
        return not self.errors

    def to_dict(self) -> dict[str, Any]:
        return {
            "filename": self.filename,
            "kind": self.kind,
            "ok": self.ok,
            "checked": self.checked,
            "errors": [i.to_dict() for i in self.errors],
            "warnings": [i.to_dict() for i in self.warnings],
        }

    def summary(self) -> str:
        if self.ok and not self.warnings:
            return f"{self.filename}：檢查通過"
        parts = []
        if self.errors:
            parts.append(f"{len(self.errors)} 個要修正")
        if self.warnings:
            parts.append(f"{len(self.warnings)} 個提醒")
        return f"{self.filename}：" + "、".join(parts)

    def repair_instruction(self) -> str:
        """給模型的修正指示。只列 error，warning 不強迫改。"""
        if not self.errors:
            return ""
        lines = [
            f"你剛才產出的「{self.filename}」有以下問題，請修正後**用同一個工具重做一次**"
            f"（檔名沿用 {Path(self.filename).stem}）：",
            "",
        ]
        for i, issue in enumerate(self.errors, 1):
            lines.append(f"{i}. {issue.message}")
            if issue.fix_hint:
                lines.append(f"   → {issue.fix_hint}")
        lines.append("")
        lines.append("只修這些問題，其他內容保持原樣，不要整份重寫成完全不同的東西。")
        return "\n".join(lines)


# ── 規則 ─────────────────────────────────────────────────────

# 療效／保證類承諾。社團是學生社團，寫這些會出事。
HEALTH_CLAIMS = re.compile(
    r"(治療|治好|根治|療效|痊癒|醫治|改善憂鬱|治癒|消除病|保證(成功|上榜|考上|錄取|進步|有效)|"
    r"一定會(好|成功|考上)|改運|開智慧|消業障|增加壽命|包治|藥效)"
)

# 對外文宣不該以宗教框架開場
RELIGIOUS_LEAD = re.compile(r"(皈依|入教|信仰|傳教|佈道|受戒|供養|法會|超渡|功德迴向)")

# 學年度／西元年
ACADEMIC_YEAR = re.compile(r"(\d{3})\s*學年度")
GREGORIAN_YEAR = re.compile(r"(20\d{2})\s*年")

# 「社長：某某」這種對今年事實的斷言
PRESIDENT_ASSERTION = re.compile(r"社長[：:]\s*([^\s，,。、|｜\n]{2,6})")
PLACEHOLDERS = {"待填", "未定", "tbd", "TBD", "unknown", "—", "-", "N/A", "n/a", "（待填）"}

# 各任務類型的必要章節
REQUIRED_SECTIONS: dict[str, tuple[str, ...]] = {
    "event_planning": ("目標", "流程"),
    "recruitment": ("目標", "管道"),
    "evaluation": ("社團", "成果"),
    "handover": ("職", "交接"),
    "finance": ("收入", "支出"),
}


def _extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    try:
        if suffix == ".docx":
            from docx import Document

            doc = Document(str(path))
            parts = [p.text for p in doc.paragraphs]
            for t in doc.tables:
                for row in t.rows:
                    parts.extend(c.text for c in row.cells)
            return "\n".join(parts)
        if suffix == ".pptx":
            from pptx import Presentation

            prs = Presentation(str(path))
            parts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        parts.append(shape.text_frame.text)
            return "\n".join(parts)
        if suffix in {".md", ".gs", ".txt", ".csv"}:
            return path.read_text(encoding="utf-8", errors="replace")
        if suffix == ".xlsx":
            from openpyxl import load_workbook

            wb = load_workbook(str(path), read_only=True, data_only=False)
            parts = []
            try:
                for ws in wb.worksheets:
                    parts.append(ws.title)
                    for row in ws.iter_rows(values_only=True):
                        parts.extend(str(c) for c in row if c is not None)
            finally:
                wb.close()
            return "\n".join(parts)
    except Exception:  # noqa: BLE001 —— 讀不到就當成沒有文字，交給呼叫端判斷
        return ""
    return ""


def _check_health_claims(text: str, report: Report) -> None:
    report.checked.append("no_health_claims")
    for m in set(HEALTH_CLAIMS.findall(text)):
        phrase = m if isinstance(m, str) else m[0]
        report.issues.append(
            Issue(
                rule="no_health_claims",
                severity="error",
                message=f"出現療效或保證類說法「{phrase}」。",
                fix_hint="禪定是自我成長的方法，不是療法。改成「幫助放鬆」「練習專注」「讓腦袋休息」這類描述。",
            )
        )


def _check_religious_framing(text: str, report: Report, external: bool) -> None:
    if not external:
        return
    report.checked.append("not_religious_recruitment")
    head = text[:600]
    hits = set(RELIGIOUS_LEAD.findall(head))
    if hits:
        report.issues.append(
            Issue(
                rule="not_religious_recruitment",
                severity="error",
                message=f"對外文宣的開頭出現宗教招募用語：{'、'.join(sorted(hits))}。",
                fix_hint="對外主軸要放在自我成長、紓壓、專注、人際、找到夥伴。禪法脈絡可以帶到，但不要從宗教講起。",
            )
        )


def _check_stale_year(text: str, report: Report) -> None:
    """歷年文件的年份被當成今年寫進產出，是最容易發生也最難察覺的錯。"""
    report.checked.append("no_stale_year_as_current")
    term = term_service.load()
    current_year = term.get("academic_year")
    if not current_year:
        return

    found = {y for y in ACADEMIC_YEAR.findall(text) if y != str(current_year)}
    # 只在「這些年份被當成本期」的情況下算錯：文件裡同時出現本學期字樣
    claims_current = any(w in text for w in ("本學期", "今年", "本年度", "這學期"))
    if found and claims_current:
        report.issues.append(
            Issue(
                rule="no_stale_year_as_current",
                severity="error",
                message=f"文件宣稱是本學期，卻出現其他學年度：{'、'.join(sorted(found))}。",
                fix_hint=f"本學期是 {current_year} 學年度。歷年資料只能當範例，不要把舊年度寫成今年。",
            )
        )
    elif found:
        report.issues.append(
            Issue(
                rule="no_stale_year_as_current",
                severity="warning",
                message=f"出現非本學期的學年度：{'、'.join(sorted(found))}。",
                fix_hint=f"確認這是刻意引用歷年資料。本學期是 {current_year} 學年度。",
            )
        )


def _check_unknown_facts_are_placeholders(text: str, report: Report) -> None:
    """當期狀態沒有社長，產出就不可以寫出一個社長姓名。"""
    report.checked.append("placeholder_for_unknown")
    term = term_service.load()
    if term.get("president"):
        return
    for name in PRESIDENT_ASSERTION.findall(text):
        if name.strip() in PLACEHOLDERS:
            continue
        report.issues.append(
            Issue(
                rule="placeholder_for_unknown",
                severity="error",
                message=f"文件寫了「社長：{name}」，但本學期設定裡還沒有社長資料。",
                fix_hint="這個欄位改成「待填」，並在回覆裡提醒使用者到「本學期設定」補上。",
            )
        )
        break


def _check_required_sections(text: str, report: Report, task_type: str) -> None:
    needed = REQUIRED_SECTIONS.get(task_type)
    if not needed:
        return
    report.checked.append("required_sections")
    missing = [s for s in needed if s not in text]
    if missing:
        report.issues.append(
            Issue(
                rule="required_sections",
                severity="error" if len(missing) == len(needed) else "warning",
                message=f"缺少必要內容：{'、'.join(missing)}。",
                fix_hint="補上這些段落。可以先用 search_knowledge 查對應的任務劇本確認完整章節。",
            )
        )


def _check_spreadsheet(path: Path, report: Report, task_type: str) -> None:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(str(path), data_only=False)
    except Exception as exc:  # noqa: BLE001
        report.issues.append(
            Issue("openable", "error", f"試算表打不開：{exc}", "重新產生一次。")
        )
        return

    report.checked += ["numeric_types", "has_header", "no_empty_sheet"]
    numeric_pat = re.compile(r"^-?\d+(\.\d+)?$")
    has_formula = False

    try:
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            if len(rows) < 2:
                report.issues.append(
                    Issue("no_empty_sheet", "warning", f"工作表「{ws.title}」只有標題沒有資料。", "補上至少一列範例。")
                )
                continue

            header = [str(c).strip() for c in rows[0] if c is not None and str(c).strip()]
            if len(header) < 2:
                report.issues.append(
                    Issue("has_header", "error", f"工作表「{ws.title}」沒有欄位標題列。", "第一列要放欄位名稱。")
                )

            for row in rows[1:]:
                for cell in row:
                    if isinstance(cell, str):
                        if cell.startswith("="):
                            has_formula = True
                        elif numeric_pat.match(cell.strip()):
                            report.issues.append(
                                Issue(
                                    "numeric_types",
                                    "warning",
                                    f"工作表「{ws.title}」有數字被存成文字（{cell}），公式會算不到。",
                                    "數字欄位直接填數字，不要加引號或前後空白。",
                                )
                            )
                            break
                else:
                    continue
                break
    finally:
        wb.close()

    if task_type == "finance":
        report.checked.append("formulas_present")
        if not has_formula:
            report.issues.append(
                Issue(
                    "formulas_present",
                    "error",
                    "經費類試算表沒有任何公式，小計與總計是手算填死的。",
                    "小計用 =B2*C2、總計用 =SUM(...)，這樣改數量才會自動重算。",
                )
            )


def verify(
    path: Path,
    *,
    task_type: str = "",
    rules: list[str] | None = None,
    external: bool = False,
) -> Report:
    """檢查一份產出。回傳 Report，呼叫端決定要不要修。"""
    path = Path(path)
    report = Report(filename=path.name, kind=path.suffix.lstrip(".").lower())

    if not path.exists():
        report.issues.append(Issue("exists", "error", "檔案不存在。", "重新產生一次。"))
        return report
    if path.stat().st_size == 0:
        report.issues.append(Issue("exists", "error", "檔案是空的。", "重新產生一次。"))
        return report

    active = set(rules or [])
    text = _extract_text(path)
    report.text_length = len(text)

    if report.kind in {"docx", "pptx", "md", "gs"} and len(text.strip()) < 40:
        report.issues.append(
            Issue("has_content", "error", "檔案幾乎沒有內容。", "把完整內容寫進去，不要只給大綱。")
        )

    if not active or "no_health_claims" in active:
        _check_health_claims(text, report)
    if not active or "not_religious_recruitment" in active:
        _check_religious_framing(text, report, external or "external_tone" in active)
    if not active or "no_stale_year_as_current" in active:
        _check_stale_year(text, report)
    if not active or "placeholder_for_unknown" in active:
        _check_unknown_facts_are_placeholders(text, report)
    if (not active or "required_sections" in active) and report.kind in {"docx", "md"}:
        _check_required_sections(text, report, task_type)
    if report.kind == "xlsx":
        _check_spreadsheet(path, report, task_type)

    return report
