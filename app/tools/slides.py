"""建立簡報（.pptx）。用於社課投影片、招生說明會、幹部交接簡報、評鑑簡報。"""

from __future__ import annotations

import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

from .base import Artifact, dated_dir, deliver, safe_filename, unique_path

PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

ACCENT = RGBColor(0x1F, 0x56, 0x73)
INK = RGBColor(0x1B, 0x24, 0x2B)
MUTED = RGBColor(0x5A, 0x6B, 0x77)
PAPER = RGBColor(0xFA, 0xFC, 0xFD)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

_SLIDE_HEAD = re.compile(r"^#{1,3}\s+(.*)$")
_BULLET = re.compile(r"^(\s*)[-*+]\s+(.*)$")


def _parse(markdown: str) -> list[dict]:
    """切成一張張投影片。`## 標題` 或 `---` 都算分頁。"""
    slides: list[dict] = []
    current: dict | None = None

    for raw in markdown.replace("\r\n", "\n").split("\n"):
        line = raw.rstrip()
        stripped = line.strip()

        if re.match(r"^\s*-{3,}\s*$", stripped):
            if current:
                slides.append(current)
                current = None
            continue

        m = _SLIDE_HEAD.match(stripped)
        if m:
            if current:
                slides.append(current)
            current = {"title": m.group(1).strip(), "bullets": [], "notes": []}
            continue

        if not stripped:
            continue

        if current is None:
            current = {"title": stripped, "bullets": [], "notes": []}
            continue

        m = _BULLET.match(line)
        if m:
            level = min(len(m.group(1)) // 2, 2)
            current["bullets"].append((level, m.group(2).strip()))
        else:
            current["bullets"].append((0, stripped))

    if current:
        slides.append(current)
    return slides


def _blank(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _rect(slide, x, y, w, h, color):
    from pptx.enum.shapes import MSO_SHAPE

    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _textbox(slide, x, y, w, h):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    return tf


def create_slides(
    filename: str,
    slides_markdown: str,
    title: str = "",
    subtitle: str = "",
    destination: str | None = None,
) -> dict:
    parsed = _parse(slides_markdown or "")
    if not parsed:
        return {"ok": False, "message": "slides_markdown 沒有可用內容。每一頁請用 `## 頁面標題` 開頭，下面接 - 項目。"}

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # 封面
    if title:
        s = _blank(prs)
        _rect(s, 0, 0, SLIDE_W, SLIDE_H, PAPER)
        _rect(s, 0, 0, Inches(0.32), SLIDE_H, ACCENT)
        tf = _textbox(s, Inches(1.1), Inches(2.5), Inches(11.0), Inches(2.2))
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = title
        run.font.size = Pt(44)
        run.font.bold = True
        run.font.color.rgb = ACCENT
        run.font.name = "微軟正黑體"
        if subtitle:
            p2 = tf.add_paragraph()
            p2.space_before = Pt(18)
            r2 = p2.add_run()
            r2.text = subtitle
            r2.font.size = Pt(20)
            r2.font.color.rgb = MUTED
            r2.font.name = "微軟正黑體"

    for item in parsed:
        s = _blank(prs)
        _rect(s, 0, 0, SLIDE_W, SLIDE_H, PAPER)
        _rect(s, Inches(0.9), Inches(0.72), Inches(0.09), Inches(0.52), ACCENT)

        tf = _textbox(s, Inches(1.16), Inches(0.6), Inches(11.3), Inches(0.9))
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = item["title"]
        run.font.size = Pt(30)
        run.font.bold = True
        run.font.color.rgb = ACCENT
        run.font.name = "微軟正黑體"

        if not item["bullets"]:
            continue

        body = _textbox(s, Inches(1.16), Inches(1.75), Inches(11.1), Inches(5.1))
        n = len(item["bullets"])
        size = 20 if n <= 6 else (17 if n <= 9 else 14)

        for idx, (level, text) in enumerate(item["bullets"]):
            para = body.paragraphs[0] if idx == 0 else body.add_paragraph()
            para.space_after = Pt(10 if n <= 8 else 6)
            para.level = level
            marker = "" if level == 0 else "· "
            run = para.add_run()
            run.text = f"{marker}{text}"
            run.font.size = Pt(size - level * 2)
            run.font.color.rgb = INK if level == 0 else MUTED
            run.font.name = "微軟正黑體"

        # 頁碼
        num = _textbox(s, SLIDE_W - Inches(1.5), SLIDE_H - Inches(0.72), Inches(0.9), Inches(0.4))
        nr = num.paragraphs[0].add_run()
        nr.text = str(prs.slides.index(s) + 1)
        nr.font.size = Pt(11)
        nr.font.color.rgb = MUTED

    name = safe_filename(filename, ".pptx")
    path = unique_path(dated_dir(), name)
    prs.save(path)

    art = Artifact(filename=path.name, summary=f"共 {len(prs.slides)} 頁。")
    return deliver(path, destination, art, mime=PPTX_MIME).to_result()


SCHEMA = {
    "type": "function",
    "function": {
        "name": "create_slides",
        "description": (
            "建立簡報（.pptx）。適用於：社課投影片、招生說明會簡報、期初茶會簡報、"
            "幹部交接簡報、社團評鑑簡報、營隊行前說明。"
        ),
        "parameters": {
            "type": "object",
            "properties": {
                "filename": {"type": "string", "description": "檔名，不用加副檔名"},
                "title": {"type": "string", "description": "封面主標題"},
                "subtitle": {"type": "string", "description": "封面副標題，例如「淡江大學領袖禪學社 · 115學年度」"},
                "slides_markdown": {
                    "type": "string",
                    "description": (
                        "投影片內容。每一頁用 `## 頁面標題` 開頭，下面接 `- 項目`。"
                        "縮排兩格代表次階層。範例：\\n"
                        "## 我們是誰\\n- 創社於民國78年\\n- 以禪定培養青年領袖\\n"
                        "## 社課在做什麼\\n- 靜心練習\\n  - 明心輪\\n- 主題分享"
                    ),
                },
                "destination": {
                    "type": "string",
                    "enum": ["local", "drive", "both"],
                    "description": "落點：local=只存本機、drive=上傳Google雲端硬碟、both=兩者",
                },
            },
            "required": ["filename", "slides_markdown"],
        },
    },
}
