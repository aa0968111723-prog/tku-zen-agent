"""建立 Word 文件（.docx）。模型寫 Markdown，這裡轉成有樣式的文件。

支援的 Markdown：標題、段落、項目符號、編號清單、表格、引言、分隔線、
粗體（**）與行內程式碼（`）。
"""

from __future__ import annotations

import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from .base import Artifact, dated_dir, deliver, safe_filename, unique_path

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

_HEADING = re.compile(r"^(#{1,5})\s+(.*)$")
_BULLET = re.compile(r"^\s*[-*+]\s+(.*)$")
_NUMBERED = re.compile(r"^\s*\d+[.)]\s+(.*)$")
_TABLE_SEP = re.compile(r"^\s*\|?[\s:|-]+\|[\s:|-]*$")
_RULE = re.compile(r"^\s*([-*_])\s*\1\s*\1[\s\-*_]*$")
_QUOTE = re.compile(r"^\s*>\s?(.*)$")
_INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")

ACCENT = RGBColor(0x1F, 0x56, 0x73)


def _add_runs(paragraph, text: str) -> None:
    """處理 **粗體** 與 `程式碼`。"""
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        else:
            paragraph.add_run(part)


def _split_row(line: str) -> list[str]:
    cells = line.strip().strip("|").split("|")
    return [c.strip() for c in cells]


def _markdown_to_docx(doc: Document, markdown: str) -> None:
    lines = markdown.replace("\r\n", "\n").split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if _RULE.match(stripped):
            doc.add_paragraph().add_run("─" * 40).font.color.rgb = RGBColor(0xC0, 0xC8, 0xCE)
            i += 1
            continue

        m = _HEADING.match(stripped)
        if m:
            level = min(len(m.group(1)), 4)
            h = doc.add_heading(level=level)
            _add_runs(h, m.group(2))
            for run in h.runs:
                run.font.color.rgb = ACCENT
            i += 1
            continue

        # 表格：目前這行有 | 且下一行是分隔線
        if "|" in stripped and i + 1 < len(lines) and _TABLE_SEP.match(lines[i + 1]):
            header = _split_row(stripped)
            body: list[list[str]] = []
            j = i + 2
            while j < len(lines) and "|" in lines[j] and lines[j].strip():
                body.append(_split_row(lines[j]))
                j += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Light Grid Accent 1"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for c, text in enumerate(header):
                cell = table.rows[0].cells[c]
                cell.text = ""
                _add_runs(cell.paragraphs[0], text)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            for row in body:
                cells = table.add_row().cells
                for c in range(len(header)):
                    cells[c].text = ""
                    _add_runs(cells[c].paragraphs[0], row[c] if c < len(row) else "")
            doc.add_paragraph()
            i = j
            continue

        m = _QUOTE.match(line)
        if m:
            p = doc.add_paragraph(style="Intense Quote")
            _add_runs(p, m.group(1))
            i += 1
            continue

        m = _BULLET.match(line)
        if m:
            indent = (len(line) - len(line.lstrip())) // 2
            style = "List Bullet" if indent == 0 else f"List Bullet {min(indent + 1, 3)}"
            try:
                p = doc.add_paragraph(style=style)
            except KeyError:
                p = doc.add_paragraph(style="List Bullet")
            _add_runs(p, m.group(1))
            i += 1
            continue

        m = _NUMBERED.match(line)
        if m:
            p = doc.add_paragraph(style="List Number")
            _add_runs(p, m.group(1))
            i += 1
            continue

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _add_runs(p, stripped)
        i += 1


def create_document(
    filename: str,
    markdown: str,
    title: str = "",
    destination: str | None = None,
    also_markdown: bool = False,
) -> dict:
    if not (markdown or "").strip():
        return {"ok": False, "message": "markdown 內容是空的，沒有東西可以寫。"}

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "微軟正黑體"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.35

    if title:
        h = doc.add_heading(level=0)
        _add_runs(h, title)
        for run in h.runs:
            run.font.color.rgb = ACCENT

    _markdown_to_docx(doc, markdown)

    name = safe_filename(filename, ".docx")
    path = unique_path(dated_dir(), name)
    doc.save(path)

    details: list[str] = []
    if also_markdown:
        md_path = path.with_suffix(".md")
        md_path.write_text(markdown, encoding="utf-8")
        details.append(f"　· 同時存了純文字版：{md_path.name}")

    words = len(re.sub(r"\s+", "", markdown))
    art = Artifact(filename=path.name, summary=f"約 {words} 字。", details=details)
    return deliver(path, destination, art, mime=DOCX_MIME).to_result()


SCHEMA = {
    "type": "function",
    "function": {
        "name": "create_document",
        "description": (
            "建立 Word 文件（.docx）。適用於：活動企劃書、招生企劃、社課教案、"
            "幹部交接手冊、社團評鑑報告、公文、演講邀請函、成果報告等。"
            "內容用 Markdown 撰寫，會自動轉成有標題階層、清單與表格樣式的正式文件。"
        ),
        "parameters": {
            "type": "object",
            "properties": {
                "filename": {"type": "string", "description": "檔名，中文可以，不用加副檔名"},
                "title": {"type": "string", "description": "文件封面大標題（可留空）"},
                "markdown": {
                    "type": "string",
                    "description": (
                        "文件正文，用 Markdown。支援 # 標題、- 項目、1. 編號、"
                        "| 表格 |、> 引言、**粗體**。請寫完整內容，不要只寫大綱。"
                    ),
                },
                "destination": {
                    "type": "string",
                    "enum": ["local", "drive", "both"],
                    "description": "落點：local=只存本機、drive=上傳Google雲端硬碟、both=兩者",
                },
                "also_markdown": {
                    "type": "boolean",
                    "description": "是否同時存一份 .md 純文字檔（方便貼到別的地方）",
                },
            },
            "required": ["filename", "markdown"],
        },
    },
}
