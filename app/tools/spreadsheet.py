"""建立試算表（.xlsx，可直接上傳成 Google 試算表）。

給模型的介面刻意用 CSV 字串而不是巢狀陣列 —— 開源模型產 CSV 的正確率
遠高於產深層 JSON。格式與型別的處理都在這邊做掉。
"""

from __future__ import annotations

import csv
import io
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .base import Artifact, dated_dir, deliver, safe_filename, unique_path

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

HEADER_FILL = PatternFill("solid", fgColor="1F5673")   # 沉靜的藍
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
ZEBRA_FILL = PatternFill("solid", fgColor="F2F6F8")
THIN = Side(style="thin", color="D3DEE5")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

_NUMBER = re.compile(r"^-?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?$")
_PERCENT = re.compile(r"^-?\d+(?:\.\d+)?%$")
_ILLEGAL_SHEET = re.compile(r"[\[\]:*?/\\]")


def _display_width(value: Any) -> int:
    """中文字寬度算 2。"""
    text = "" if value is None else str(value)
    return sum(2 if ord(ch) > 0x2E80 else 1 for ch in text)


def _typed(raw: str) -> Any:
    s = (raw or "").strip()
    if not s:
        return None
    if s.startswith("="):          # 公式原樣寫入
        return s
    if _NUMBER.match(s):
        n = float(s.replace(",", ""))
        return int(n) if n.is_integer() else n
    if _PERCENT.match(s):
        return float(s[:-1]) / 100
    return s


def _parse_csv(text: str) -> list[list[str]]:
    text = (text or "").strip()
    if not text:
        return []
    # 模型偶爾會包 code fence
    fence = re.match(r"^```(?:csv|tsv)?\s*(.+?)\s*```$", text, re.S)
    if fence:
        text = fence.group(1)
    sample = text.splitlines()[0] if text else ""
    delimiter = "\t" if sample.count("\t") > sample.count(",") else ","
    return [row for row in csv.reader(io.StringIO(text), delimiter=delimiter) if any(c.strip() for c in row)]


def _normalize_sheets(sheets: Any) -> list[dict[str, str]]:
    """把模型可能給的各種形狀，收斂成 [{'name':..., 'csv':...}]。"""
    if isinstance(sheets, str):
        try:
            import json

            sheets = json.loads(sheets)
        except Exception:  # noqa: BLE001 —— 純字串就當成單一工作表的 CSV
            return [{"name": "工作表1", "csv": sheets}]
    if isinstance(sheets, dict):
        # 可能是 {"工作表名": "csv..."} 或單一 {"name":..., "csv":...}
        if "csv" in sheets or "rows" in sheets or "data" in sheets:
            sheets = [sheets]
        else:
            return [{"name": str(k), "csv": str(v)} for k, v in sheets.items()]
    if not isinstance(sheets, list):
        return []

    out: list[dict[str, str]] = []
    for i, item in enumerate(sheets, 1):
        if isinstance(item, str):
            out.append({"name": f"工作表{i}", "csv": item})
            continue
        if not isinstance(item, dict):
            continue
        name = str(item.get("name") or item.get("title") or item.get("sheet") or f"工作表{i}")
        body = item.get("csv") or item.get("rows") or item.get("data") or item.get("content") or ""
        if isinstance(body, list):
            body = "\n".join(
                ",".join('"' + str(c).replace('"', '""') + '"' for c in row) if isinstance(row, list) else str(row)
                for row in body
            )
        out.append({"name": name, "csv": str(body)})
    return out


def _write_sheet(ws: Worksheet, rows: list[list[str]], freeze_header: bool) -> None:
    if not rows:
        ws["A1"] = "（沒有資料）"
        return

    for r_i, row in enumerate(rows, start=1):
        for c_i, cell in enumerate(row, start=1):
            ws.cell(row=r_i, column=c_i, value=_typed(cell))

    n_cols = max(len(r) for r in rows)

    for c_i in range(1, n_cols + 1):
        head = ws.cell(row=1, column=c_i)
        head.fill = HEADER_FILL
        head.font = HEADER_FONT
        head.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r_i in range(1, len(rows) + 1):
        for c_i in range(1, n_cols + 1):
            cell = ws.cell(row=r_i, column=c_i)
            cell.border = BORDER
            if r_i == 1:
                continue
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if r_i % 2 == 0:
                cell.fill = ZEBRA_FILL

    for c_i in range(1, n_cols + 1):
        widest = max(
            (_display_width(ws.cell(row=r, column=c_i).value) for r in range(1, len(rows) + 1)),
            default=8,
        )
        ws.column_dimensions[get_column_letter(c_i)].width = max(9, min(52, widest + 3))

    ws.row_dimensions[1].height = 26
    if freeze_header:
        ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(n_cols)}{len(rows)}"


def create_spreadsheet(
    filename: str,
    sheets: Any,
    destination: str | None = None,
    summary: str = "",
) -> dict:
    parsed = _normalize_sheets(sheets)
    if not parsed:
        return {"ok": False, "message": "sheets 參數是空的或格式不對。請傳一個陣列，每個元素是 {\"name\": 工作表名稱, \"csv\": CSV內容}。"}

    wb = Workbook()
    wb.remove(wb.active)

    used: set[str] = set()
    total_rows = 0
    details: list[str] = []

    for item in parsed:
        rows = _parse_csv(item["csv"])
        title = _ILLEGAL_SHEET.sub("_", item["name"]).strip()[:31] or "工作表"
        base_title = title
        n = 2
        while title in used:
            title = f"{base_title[:28]}_{n}"
            n += 1
        used.add(title)

        ws = wb.create_sheet(title=title)
        _write_sheet(ws, rows, freeze_header=True)
        body_rows = max(0, len(rows) - 1)
        total_rows += body_rows
        details.append(f"　· {title}：{body_rows} 筆資料、{max((len(r) for r in rows), default=0)} 欄")

    name = safe_filename(filename, ".xlsx")
    path = unique_path(dated_dir(), name)
    wb.save(path)

    art = Artifact(
        filename=path.name,
        summary=summary or f"共 {len(parsed)} 個工作表、{total_rows} 筆資料。",
        details=details,
    )
    return deliver(path, destination, art, mime=XLSX_MIME).to_result()


SCHEMA = {
    "type": "function",
    "function": {
        "name": "create_spreadsheet",
        "description": (
            "建立 Excel 試算表（.xlsx），會自動排版（標題列上色、凍結、自動欄寬、篩選器）。"
            "適用於：社課排程表、招生名單、幹部分工表、經費預算表、器材清單、報名統計、社團評鑑資料表等。"
            "數字會自動存成數值型別，所以可以用 =SUM() 之類的公式。"
        ),
        "parameters": {
            "type": "object",
            "properties": {
                "filename": {
                    "type": "string",
                    "description": "檔名，中文可以，不用加副檔名。例如「115學年度上學期社課排程」",
                },
                "sheets": {
                    "type": "array",
                    "description": "工作表清單。第一列必須是欄位標題。",
                    "items": {
                        "type": "object",
                        "properties": {
                            "name": {"type": "string", "description": "工作表名稱，例如「社課排程」"},
                            "csv": {
                                "type": "string",
                                "description": (
                                    "該工作表的 CSV 內容，第一行是欄位標題。"
                                    "含逗號的欄位請用雙引號包起來。"
                                    "例如：週次,日期,主題,帶課者\\n第1週,9/15,期初茶會,社長"
                                ),
                            },
                        },
                        "required": ["name", "csv"],
                    },
                },
                "destination": {
                    "type": "string",
                    "enum": ["local", "drive", "both"],
                    "description": "落點：local=只存本機、drive=上傳Google雲端硬碟、both=兩者。沒指定就用使用者的預設值。",
                },
                "summary": {"type": "string", "description": "一句話說明這份表是做什麼用的"},
            },
            "required": ["filename", "sheets"],
        },
    },
}
