"""安全讀取既有產出，讓「把上一份改成簡報」使用真正內容而非只看片名。"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from .. import config
from ..services import context as ctx_mod
from ..services.session_store import get_store


def _extract(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".md", ".txt", ".csv", ".tsv", ".gs", ".json"}:
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".docx":
        from docx import Document

        doc = Document(str(path))
        parts = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]
        for table in doc.tables:
            parts.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(parts)
    if suffix == ".pptx":
        from pptx import Presentation

        presentation = Presentation(str(path))
        parts: list[str] = []
        for index, slide in enumerate(presentation.slides, 1):
            parts.append(f"## 投影片 {index}")
            for shape in slide.shapes:
                text = getattr(shape, "text", "")
                if text and text.strip():
                    parts.append(text.strip())
        return "\n".join(parts)
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, data_only=False, read_only=True)
        parts: list[str] = []
        try:
            for sheet in workbook.worksheets:
                parts.append(f"## 工作表：{sheet.title}")
                for row in sheet.iter_rows(max_row=500, values_only=True):
                    if any(value is not None for value in row):
                        parts.append(" | ".join("" if value is None else str(value) for value in row))
        finally:
            workbook.close()
        return "\n".join(parts)
    raise ValueError(f"目前不支援讀取 {suffix or '未知格式'}；可先下載後貼上文字，或轉成 docx／pptx／xlsx／md。")


def read_artifact(artifact_id: str = "", max_chars: int = 24000) -> dict[str, Any]:
    ctx = ctx_mod.current()
    if ctx is None:
        raise RuntimeError("讀取產出需要登入後的工作情境")
    store = get_store()
    record = store.get_artifact(artifact_id, ctx.user_id) if artifact_id else None
    if record is None and not artifact_id:
        latest = store.list_artifacts(ctx.user_id, project_id=ctx.project_id, limit=1)
        record = latest[0] if latest else None
    if record is None:
        return {
            "ok": False, "code": "artifact_not_found", "message": "找不到指定產出或上一份產出。",
            "alternative": "可以重新上傳／產生檔案，或直接貼上要轉換的內容。",
        }
    path = Path(record.local_path).resolve()
    root = config.OUTPUT_DIR.resolve()
    if not path.is_relative_to(root) or not path.is_file():
        return {
            "ok": False, "code": "artifact_unavailable", "message": "這份產出的本機檔案已不存在。",
            "alternative": "若有雲端版本可先下載回來，或貼上原文後接續修改。",
        }
    try:
        content = _extract(path)
    except (OSError, ValueError) as exc:
        return {
            "ok": False, "code": "artifact_unreadable", "message": str(exc),
            "alternative": "可下載後另存為 docx、pptx、xlsx 或 Markdown，再重新匯入。",
        }
    limit = max(1000, min(int(max_chars or 24000), 50000))
    truncated = len(content) > limit
    content = content[:limit]
    return {
        "ok": True,
        "message": f"已讀取 {record.filename} 第 {record.version} 版" + ("（內容較長，已截取前段）" if truncated else ""),
        "artifact": record.public(),
        "content": content,
        "truncated": truncated,
    }


SCHEMA = {
    "type": "function",
    "function": {
        "name": "read_artifact",
        "description": "讀取使用者自己的既有產出內容。轉成簡報、Reels、網宣或修改上一份前必須先呼叫。",
        "parameters": {
            "type": "object",
            "properties": {
                "artifact_id": {"type": "string", "description": "上一份產出的 artifact_id；留空則讀同 project 最新產出"},
                "max_chars": {"type": "integer", "minimum": 1000, "maximum": 50000},
            },
        },
    },
}
