"""Provider boundaries for duigao Asset Intelligence.

The duigao edge function owns room membership, Storage signing and the
external-AI policy.  This module only understands an already-authorised,
short-lived source URL and returns bounded structured data.  Provider-specific
implementations stay behind small protocols so a self-hosted video or OCR
service can replace FAL without changing the request contract.
"""

from __future__ import annotations

import json
from dataclasses import dataclass, field
from io import BytesIO
from typing import Any, Protocol

import httpx

from .. import config


class AssetProviderError(RuntimeError):
    """An expected provider failure with a safe, machine-readable code."""

    def __init__(self, message: str, *, code: str = "PROVIDER_UNAVAILABLE") -> None:
        super().__init__(message)
        self.code = code


@dataclass(frozen=True)
class AIProviderPolicy:
    """Capabilities/data policy declared by one multimodal provider."""

    name: str
    supports_image: bool = False
    supports_video: bool = False
    supports_audio: bool = False
    supports_documents: bool = False
    is_external: bool = True
    data_retention_policy: str = "request-only"
    allowed_for_private_assets: bool = False

    def supports(self, asset_type: str) -> bool:
        return {
            "image": self.supports_image,
            "video": self.supports_video,
            "audio": self.supports_audio,
            "document": self.supports_documents,
            "plan": self.supports_documents,
        }.get(asset_type, False)


class EmbeddingProvider(Protocol):
    """Optional vector provider; lexical retrieval remains the safe fallback."""

    @property
    def enabled(self) -> bool: ...

    async def embed(self, texts: list[str]) -> list[list[float]]: ...


class DisabledEmbeddingProvider:
    enabled = False

    async def embed(self, texts: list[str]) -> list[list[float]]:
        return []


def approved_provider_policy() -> AIProviderPolicy:
    external = config.DUIGAO_ANALYSIS_PROVIDER_EXTERNAL
    return AIProviderPolicy(
        name="tku-zen-agent",
        supports_image=True,
        supports_video=bool(config.FAL_VIDEO_ENDPOINT),
        supports_audio=bool(config.AUDIO_TRANSCRIPTION_ENDPOINT),
        supports_documents=True,
        is_external=external,
        data_retention_policy="provider-defined" if external else "request-only",
        allowed_for_private_assets=not external,
    )


@dataclass(frozen=True)
class DocumentPage:
    page: int | None
    text: str
    heading: str | None = None


@dataclass(frozen=True)
class DocumentExtraction:
    text: str
    pages: list[DocumentPage] = field(default_factory=list)
    mime_type: str = ""


class VideoUnderstandingProvider(Protocol):
    """Turn an approved video URL/keyframe set into temporal segments."""

    async def understand(
        self,
        *,
        source_url: str,
        title: str,
        mime_type: str,
        keyframes: list[dict[str, Any]],
        duration_seconds: float | None,
    ) -> dict[str, Any]: ...


class DocumentUnderstandingProvider(Protocol):
    """Extract searchable text and page boundaries without prompting a model."""

    async def extract(
        self,
        *,
        source_url: str | None,
        text_content: str,
        title: str,
        mime_type: str,
    ) -> DocumentExtraction: ...


class AudioUnderstandingProvider(Protocol):
    """Return timestamped transcript data when an approved service is configured."""

    async def transcribe(
        self,
        *,
        source_url: str,
        title: str,
        mime_type: str,
    ) -> dict[str, Any]: ...


def _content_type(mime_type: str, url: str = "") -> str:
    value = (mime_type or "").lower().split(";", 1)[0].strip()
    if value:
        return value
    lower = url.lower().split("?", 1)[0]
    if lower.endswith(".pdf"):
        return "application/pdf"
    if lower.endswith(".docx"):
        return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if lower.endswith(('.md', '.markdown')):
        return "text/markdown"
    return "text/plain"


def _bounded_text(value: str, limit: int = 50_000) -> str:
    return value.replace("\x00", "").strip()[:limit]


class LocalDocumentUnderstandingProvider:
    """Parse common office/document formats using tku's existing dependencies."""

    MAX_BYTES = 25 * 1024 * 1024

    async def extract(
        self,
        *,
        source_url: str | None,
        text_content: str,
        title: str,
        mime_type: str,
    ) -> DocumentExtraction:
        if text_content.strip():
            text = _bounded_text(text_content)
            return DocumentExtraction(text=text, pages=[DocumentPage(page=None, text=text)], mime_type=_content_type(mime_type, title))
        if not source_url:
            raise AssetProviderError("文件沒有可抽取的文字或暫時檔案網址。", code="UNSUPPORTED_FORMAT")
        try:
            async with httpx.AsyncClient(timeout=httpx.Timeout(45.0, connect=10.0), follow_redirects=False) as client:
                response = await client.get(source_url)
        except httpx.TimeoutException as exc:
            raise AssetProviderError("文件下載逾時。", code="ANALYSIS_TIMEOUT") from exc
        except httpx.HTTPError as exc:
            raise AssetProviderError("文件暫時無法下載。", code="PROVIDER_UNAVAILABLE") from exc
        if response.status_code >= 400:
            raise AssetProviderError("文件暫時無法下載。", code="PROVIDER_UNAVAILABLE")
        if len(response.content) > self.MAX_BYTES:
            raise AssetProviderError("文件超過可分析大小。", code="FILE_TOO_LARGE")
        kind = _content_type(mime_type, title)
        data = response.content
        if kind in {"text/plain", "text/markdown", "text/csv"}:
            text = _bounded_text(data.decode("utf-8", errors="replace"))
            return DocumentExtraction(text=text, pages=[DocumentPage(page=None, text=text)], mime_type=kind)
        if kind == "application/pdf" or title.lower().endswith(".pdf"):
            try:
                from pypdf import PdfReader

                pages = [
                    DocumentPage(page=index, text=_bounded_text(page.extract_text() or ""))
                    for index, page in enumerate(PdfReader(BytesIO(data)).pages, start=1)
                ]
            except Exception as exc:  # parser details must not cross the API boundary
                raise AssetProviderError("PDF 文字抽取失敗。", code="OCR_FAILED") from exc
            return DocumentExtraction(
                text=_bounded_text("\n\n".join(page.text for page in pages)),
                pages=pages,
                mime_type="application/pdf",
            )
        if "wordprocessingml.document" in kind or title.lower().endswith(".docx"):
            try:
                from docx import Document

                document = Document(BytesIO(data))
                paragraphs = [_bounded_text(item.text, 5000) for item in document.paragraphs if item.text.strip()]
                table_lines: list[str] = []
                for table in document.tables:
                    for row in table.rows:
                        values = [_bounded_text(cell.text, 1000) for cell in row.cells]
                        if any(values):
                            table_lines.append(" | ".join(values))
                text = _bounded_text("\n".join([*paragraphs, *table_lines]))
            except Exception as exc:
                raise AssetProviderError("DOCX 文字抽取失敗。", code="OCR_FAILED") from exc
            return DocumentExtraction(text=text, pages=[DocumentPage(page=None, text=text)], mime_type=kind)
        raise AssetProviderError("目前不支援這種文件格式。", code="UNSUPPORTED_FORMAT")


class ConfiguredVideoUnderstandingProvider:
    """Call an explicitly configured temporal provider; never pretend a poster is a timeline."""

    def _endpoint(self) -> str:
        return config.FAL_VIDEO_ENDPOINT.strip()

    async def understand(
        self,
        *,
        source_url: str,
        title: str,
        mime_type: str,
        keyframes: list[dict[str, Any]],
        duration_seconds: float | None,
    ) -> dict[str, Any]:
        endpoint = self._endpoint()
        if not endpoint:
            raise AssetProviderError("尚未設定影片時間理解 provider。", code="PROVIDER_UNAVAILABLE")
        headers = {"Accept": "application/json", "Content-Type": "application/json"}
        if config.FAL_KEY:
            headers["Authorization"] = f"Key {config.FAL_KEY}"
        payload = {
            "video_url": source_url,
            "title": title,
            "mime_type": mime_type,
            "duration_seconds": duration_seconds,
            "keyframes": [
                {
                    "image_url": item.get("imageUrl") or item.get("image_url"),
                    "start_seconds": item.get("startSeconds") or item.get("start_seconds"),
                    "end_seconds": item.get("endSeconds") or item.get("end_seconds"),
                }
                for item in keyframes[:12]
            ],
            "prompt": "請回傳繁體中文 JSON，建立不重疊的時間片段 summary、transcript、topics、detectedText、sceneType、confidence。只能描述影片實際可見或可聽內容，不要猜測真人身分。",
        }
        try:
            async with httpx.AsyncClient(timeout=httpx.Timeout(120.0, connect=15.0), follow_redirects=False) as client:
                response = await client.post(endpoint, headers=headers, json=payload)
        except httpx.TimeoutException as exc:
            raise AssetProviderError("影片理解逾時。", code="ANALYSIS_TIMEOUT") from exc
        except httpx.HTTPError as exc:
            raise AssetProviderError("影片理解 provider 無法連線。", code="PROVIDER_UNAVAILABLE") from exc
        if response.status_code >= 400:
            raise AssetProviderError("影片理解 provider 回應失敗。", code="PROVIDER_UNAVAILABLE")
        try:
            data = response.json()
        except ValueError as exc:
            raise AssetProviderError("影片理解 provider 回傳格式錯誤。", code="PROVIDER_UNAVAILABLE") from exc
        if not isinstance(data, dict):
            raise AssetProviderError("影片理解 provider 回傳格式錯誤。", code="PROVIDER_UNAVAILABLE")
        output = data.get("output", data)
        if isinstance(output, str):
            try:
                output = json.loads(output)
            except json.JSONDecodeError:
                output = {"summary": output}
        return output if isinstance(output, dict) else {"summary": str(output)}


class ConfiguredAudioUnderstandingProvider:
    """Optional timestamped transcription adapter; disabled unless configured."""

    async def transcribe(self, *, source_url: str, title: str, mime_type: str) -> dict[str, Any]:
        endpoint = config.AUDIO_TRANSCRIPTION_ENDPOINT.strip()
        if not endpoint:
            raise AssetProviderError("尚未設定音訊逐字稿 provider。", code="TRANSCRIPT_FAILED")
        headers = {"Accept": "application/json", "Content-Type": "application/json"}
        if config.FAL_KEY:
            headers["Authorization"] = f"Key {config.FAL_KEY}"
        try:
            async with httpx.AsyncClient(timeout=httpx.Timeout(120.0, connect=15.0), follow_redirects=False) as client:
                response = await client.post(endpoint, headers=headers, json={"audio_url": source_url, "title": title, "mime_type": mime_type})
        except httpx.TimeoutException as exc:
            raise AssetProviderError("音訊逐字稿逾時。", code="TRANSCRIPT_FAILED") from exc
        except httpx.HTTPError as exc:
            raise AssetProviderError("音訊逐字稿 provider 無法連線。", code="PROVIDER_UNAVAILABLE") from exc
        if response.status_code >= 400:
            raise AssetProviderError("音訊逐字稿 provider 回應失敗。", code="TRANSCRIPT_FAILED")
        try:
            data = response.json()
        except ValueError as exc:
            raise AssetProviderError("音訊逐字稿格式錯誤。", code="TRANSCRIPT_FAILED") from exc
        return data if isinstance(data, dict) else {"summary": str(data)}
