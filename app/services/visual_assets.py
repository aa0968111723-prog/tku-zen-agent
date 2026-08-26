"""可追溯、非破壞式的視覺資產資料庫。

這個模組刻意和聊天 orchestrator 解耦，但共用同一個 SQLite 與 user_id：
既有對話、知識庫與產檔流程不需要改寫。所有 AI 結果都是 observation，必須
帶來源、信心與狀態；原圖路徑、雜湊與上傳者永遠不接受 PATCH。
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import re
import secrets
import sqlite3
import threading
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath
from typing import Any, Iterable

from PIL import Image, ImageFilter, ImageOps, ImageStat, UnidentifiedImageError
from docx import Document
from pypdf import PdfReader
from pptx import Presentation

from .. import config
from ..rag.hybrid import reciprocal_rank_fusion
from .visual_migrations import apply_visual_migrations


def now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def new_id(prefix: str) -> str:
    return f"{prefix}_{secrets.token_urlsafe(12)}"


def dumps(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def loads(value: str | None, default: Any) -> Any:
    try:
        parsed = json.loads(value or "")
        return parsed
    except (TypeError, ValueError):
        return default


SCHEMA = """
PRAGMA journal_mode=WAL;
PRAGMA foreign_keys=ON;

CREATE TABLE IF NOT EXISTS visual_schools (
    id TEXT PRIMARY KEY,
    canonical_name TEXT NOT NULL UNIQUE,
    aliases TEXT NOT NULL DEFAULT '[]',
    created_at TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS visual_clubs (
    id TEXT PRIMARY KEY,
    school_id TEXT NOT NULL REFERENCES visual_schools(id),
    name TEXT NOT NULL,
    aliases TEXT NOT NULL DEFAULT '[]',
    logo_asset_id TEXT NOT NULL DEFAULT '',
    social_accounts TEXT NOT NULL DEFAULT '{}',
    official_website TEXT NOT NULL DEFAULT '',
    status TEXT NOT NULL DEFAULT 'pending_review',
    source TEXT NOT NULL DEFAULT '{}',
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL,
    UNIQUE(school_id, name)
);

CREATE TABLE IF NOT EXISTS visual_people (
    id TEXT PRIMARY KEY,
    school_id TEXT REFERENCES visual_schools(id),
    club_id TEXT REFERENCES visual_clubs(id),
    name TEXT NOT NULL,
    nickname TEXT NOT NULL DEFAULT '',
    role TEXT NOT NULL DEFAULT '',
    status TEXT NOT NULL DEFAULT 'pending_review',
    source TEXT NOT NULL DEFAULT '{}',
    confidence REAL NOT NULL DEFAULT 0,
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_visual_people_name ON visual_people(name, school_id);

CREATE TABLE IF NOT EXISTS visual_scenes (
    id TEXT PRIMARY KEY,
    name TEXT NOT NULL UNIQUE,
    category TEXT NOT NULL DEFAULT '',
    created_at TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS visual_events (
    id TEXT PRIMARY KEY,
    school_id TEXT REFERENCES visual_schools(id),
    club_id TEXT REFERENCES visual_clubs(id),
    name TEXT NOT NULL,
    event_type TEXT NOT NULL DEFAULT '',
    event_date TEXT NOT NULL DEFAULT '',
    start_time TEXT NOT NULL DEFAULT '',
    end_time TEXT NOT NULL DEFAULT '',
    location TEXT NOT NULL DEFAULT '',
    copy_text TEXT NOT NULL DEFAULT '',
    status TEXT NOT NULL DEFAULT 'pending_review',
    evidence TEXT NOT NULL DEFAULT '[]',
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_visual_events_school_date ON visual_events(school_id, event_date);

CREATE TABLE IF NOT EXISTS visual_assets (
    id TEXT PRIMARY KEY,
    user_id TEXT NOT NULL REFERENCES users(id),
    original_filename TEXT NOT NULL,
    storage_path TEXT NOT NULL,
    thumbnail_path TEXT NOT NULL,
    mime_type TEXT NOT NULL,
    extension TEXT NOT NULL,
    width INTEGER NOT NULL,
    height INTEGER NOT NULL,
    orientation TEXT NOT NULL,
    uploaded_at TEXT NOT NULL,
    source TEXT NOT NULL DEFAULT '',
    source_metadata TEXT NOT NULL DEFAULT '{}',
    exif_date TEXT NOT NULL DEFAULT '',
    file_created_date TEXT NOT NULL DEFAULT '',
    sha256 TEXT NOT NULL,
    perceptual_hash TEXT NOT NULL,
    color_embedding TEXT NOT NULL DEFAULT '[]',
    semantic_embedding TEXT NOT NULL DEFAULT '[]',
    privacy TEXT NOT NULL DEFAULT 'private',
    commercial_use TEXT NOT NULL DEFAULT 'unknown',
    analysis_status TEXT NOT NULL DEFAULT 'local_complete',
    quality_score REAL NOT NULL DEFAULT 0,
    blur_score REAL NOT NULL DEFAULT 0,
    brightness_score REAL NOT NULL DEFAULT 0,
    suitability TEXT NOT NULL DEFAULT '{}',
    ocr_text TEXT NOT NULL DEFAULT '',
    analysis TEXT NOT NULL DEFAULT '{}',
    review_status TEXT NOT NULL DEFAULT 'pending_review',
    school_id TEXT REFERENCES visual_schools(id),
    club_id TEXT REFERENCES visual_clubs(id),
    duplicate_of TEXT REFERENCES visual_assets(id),
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_visual_assets_user_time ON visual_assets(user_id, uploaded_at DESC);
CREATE INDEX IF NOT EXISTS idx_visual_assets_sha ON visual_assets(sha256);
CREATE INDEX IF NOT EXISTS idx_visual_assets_school ON visual_assets(school_id, uploaded_at DESC);

CREATE TABLE IF NOT EXISTS visual_observations (
    id TEXT PRIMARY KEY,
    asset_id TEXT NOT NULL REFERENCES visual_assets(id) ON DELETE CASCADE,
    entity_type TEXT NOT NULL,
    entity_id TEXT NOT NULL DEFAULT '',
    label TEXT NOT NULL DEFAULT '',
    status TEXT NOT NULL DEFAULT 'pending_review',
    confidence REAL NOT NULL DEFAULT 0,
    source TEXT NOT NULL,
    evidence TEXT NOT NULL DEFAULT '{}',
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_visual_obs_asset ON visual_observations(asset_id, entity_type, status);
CREATE INDEX IF NOT EXISTS idx_visual_obs_entity ON visual_observations(entity_type, entity_id);

CREATE TABLE IF NOT EXISTS visual_date_candidates (
    id TEXT PRIMARY KEY,
    asset_id TEXT NOT NULL REFERENCES visual_assets(id) ON DELETE CASCADE,
    value TEXT NOT NULL,
    precision TEXT NOT NULL DEFAULT 'date',
    source TEXT NOT NULL,
    confidence REAL NOT NULL DEFAULT 0,
    status TEXT NOT NULL DEFAULT 'probable',
    evidence TEXT NOT NULL DEFAULT '',
    created_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_visual_dates_asset ON visual_date_candidates(asset_id, value);

CREATE TABLE IF NOT EXISTS visual_person_references (
    person_id TEXT NOT NULL REFERENCES visual_people(id),
    asset_id TEXT NOT NULL REFERENCES visual_assets(id),
    status TEXT NOT NULL DEFAULT 'verified',
    created_at TEXT NOT NULL,
    PRIMARY KEY(person_id, asset_id)
);

CREATE TABLE IF NOT EXISTS visual_analysis_jobs (
    id TEXT PRIMARY KEY,
    asset_id TEXT NOT NULL REFERENCES visual_assets(id),
    status TEXT NOT NULL,
    stage TEXT NOT NULL DEFAULT '',
    progress INTEGER NOT NULL DEFAULT 0,
    error_code TEXT NOT NULL DEFAULT '',
    error_message TEXT NOT NULL DEFAULT '',
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS visual_learning_events (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    user_id TEXT NOT NULL,
    event_type TEXT NOT NULL,
    query_text TEXT NOT NULL DEFAULT '',
    asset_id TEXT NOT NULL DEFAULT '',
    payload TEXT NOT NULL DEFAULT '{}',
    outcome TEXT NOT NULL DEFAULT '',
    created_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_visual_learning_type ON visual_learning_events(event_type, created_at DESC);

CREATE TABLE IF NOT EXISTS visual_corrections (
    id TEXT PRIMARY KEY,
    user_id TEXT NOT NULL,
    asset_id TEXT NOT NULL REFERENCES visual_assets(id),
    field_name TEXT NOT NULL,
    old_value TEXT NOT NULL DEFAULT '{}',
    new_value TEXT NOT NULL DEFAULT '{}',
    reason TEXT NOT NULL DEFAULT '',
    approval_status TEXT NOT NULL DEFAULT 'reviewed',
    created_at TEXT NOT NULL,
    approved_at TEXT NOT NULL DEFAULT '',
    approved_by TEXT NOT NULL DEFAULT ''
);
CREATE INDEX IF NOT EXISTS idx_visual_corrections_created ON visual_corrections(created_at DESC);

CREATE TABLE IF NOT EXISTS visual_collections (
    id TEXT PRIMARY KEY,
    user_id TEXT NOT NULL REFERENCES users(id),
    name TEXT NOT NULL,
    kind TEXT NOT NULL,
    status TEXT NOT NULL DEFAULT 'active',
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_visual_collections_user ON visual_collections(user_id,kind,updated_at DESC);

CREATE TABLE IF NOT EXISTS visual_collection_items (
    collection_id TEXT NOT NULL REFERENCES visual_collections(id) ON DELETE CASCADE,
    asset_id TEXT NOT NULL REFERENCES visual_assets(id),
    position INTEGER NOT NULL DEFAULT 0,
    note TEXT NOT NULL DEFAULT '',
    rendition TEXT NOT NULL DEFAULT 'original',
    created_at TEXT NOT NULL,
    PRIMARY KEY(collection_id,asset_id)
);
"""


SCENE_NAMES = (
    "校園", "教室", "禪堂", "舞台", "報到區", "戶外", "茶會", "講座",
    "社課", "聚餐", "活動現場", "海報", "文件",
)
IMAGE_MIME = {
    "image/jpeg": ".jpg", "image/png": ".png", "image/webp": ".webp",
    "image/gif": ".gif", "image/avif": ".avif", "image/jfif": ".jfif",
}
VIDEO_MIME = {"video/mp4": ".mp4", "video/quicktime": ".mov", "video/webm": ".webm"}
DOCUMENT_MIME = {
    "application/pdf": ".pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": ".pptx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
    "text/x-google-apps-script": ".gs",
    "text/plain": ".txt", "text/markdown": ".md", "text/csv": ".csv",
}
ALLOWED_MIME = {**IMAGE_MIME, **VIDEO_MIME, **DOCUMENT_MIME}
STATUS_VALUES = {"verified", "probable", "pending_review", "conflicted", "failed"}


class VisualAssetError(ValueError):
    def __init__(self, message: str, *, code: str = "invalid_visual_asset") -> None:
        super().__init__(message)
        self.code = code


@dataclass
class ImageInspection:
    width: int
    height: int
    orientation: str
    exif_date: str
    sha256: str
    perceptual_hash: str
    color_embedding: list[float]
    blur_score: float
    brightness_score: float
    quality_score: float
    suitability: dict[str, Any]
    thumbnail: bytes
    asset_type: str = "image"
    extracted_text: str = ""


def _variance(values: Iterable[float]) -> float:
    seq = list(values)
    if not seq:
        return 0.0
    mean = sum(seq) / len(seq)
    return sum((v - mean) ** 2 for v in seq) / len(seq)


def _image_hash(image: Image.Image) -> str:
    small = ImageOps.grayscale(image).resize((9, 8), Image.Resampling.LANCZOS)
    px = list(small.getdata())
    bits = [px[y * 9 + x] > px[y * 9 + x + 1] for y in range(8) for x in range(8)]
    return f"{sum((1 << i) for i, bit in enumerate(bits) if bit):016x}"


def _color_embedding(image: Image.Image) -> list[float]:
    rgb = image.convert("RGB").resize((96, 96), Image.Resampling.BILINEAR)
    histogram = rgb.histogram()
    bins: list[float] = []
    for channel in range(3):
        raw = histogram[channel * 256:(channel + 1) * 256]
        bins.extend(sum(raw[i:i + 16]) for i in range(0, 256, 16))
    norm = math.sqrt(sum(v * v for v in bins)) or 1.0
    return [round(v / norm, 7) for v in bins]


def inspect_image(content: bytes, mime_type: str) -> ImageInspection:
    if mime_type not in IMAGE_MIME:
        raise VisualAssetError("只支援 JPEG、PNG、WebP、GIF、AVIF、JFIF 圖片", code="unsupported_media_type")
    if not content:
        raise VisualAssetError("圖片內容是空的", code="empty_file")
    try:
        with Image.open(io.BytesIO(content)) as probe:
            probe.verify()
        with Image.open(io.BytesIO(content)) as raw:
            exif = raw.getexif()
            exif_date = str(exif.get(36867) or exif.get(306) or "").strip()
            image = ImageOps.exif_transpose(raw).convert("RGB")
    except (UnidentifiedImageError, OSError, ValueError) as exc:
        raise VisualAssetError("檔案不是可讀取的圖片或內容已損壞", code="invalid_image") from exc
    width, height = image.size
    if width < 16 or height < 16 or width * height > 120_000_000:
        raise VisualAssetError("圖片尺寸不在允許範圍", code="invalid_dimensions")
    orientation = "square" if abs(width - height) / max(width, height) < 0.04 else ("landscape" if width > height else "portrait")
    gray = ImageOps.grayscale(image).resize((min(width, 640), max(1, round(height * min(width, 640) / width))))
    brightness = float(ImageStat.Stat(gray).mean[0])
    edge = gray.filter(ImageFilter.FIND_EDGES)
    sharpness = math.sqrt(_variance(edge.getdata()))
    blur_score = max(0.0, min(100.0, sharpness * 6.5))
    brightness_score = max(0.0, min(100.0, 100.0 - abs(brightness - 132.0) * 0.85))
    resolution_score = min(100.0, math.sqrt(width * height) / 18.0)
    quality = round(blur_score * 0.45 + brightness_score * 0.25 + resolution_score * 0.30, 1)
    ratio = width / height
    suitability = {
        "poster": bool(quality >= 62 and height >= 1080),
        "video": bool(quality >= 58 and width >= 1280),
        "instagram_cover": bool(quality >= 68 and 0.75 <= ratio <= 1.91),
        "reasons": [
            *( ["可能模糊"] if blur_score < 40 else [] ),
            *( ["可能過暗"] if brightness < 55 else [] ),
            *( ["解析度偏低"] if width * height < 900_000 else [] ),
        ],
    }
    thumb = image.copy()
    thumb.thumbnail((640, 640), Image.Resampling.LANCZOS)
    thumb_io = io.BytesIO()
    thumb.save(thumb_io, format="JPEG", quality=84, optimize=True)
    return ImageInspection(
        width=width, height=height, orientation=orientation, exif_date=exif_date,
        sha256=hashlib.sha256(content).hexdigest(), perceptual_hash=_image_hash(image),
        color_embedding=_color_embedding(image), blur_score=round(blur_score, 1),
        brightness_score=round(brightness_score, 1), quality_score=quality,
        suitability=suitability, thumbnail=thumb_io.getvalue(),
    )


def _placeholder_thumbnail(title: str, subtitle: str) -> bytes:
    canvas = Image.new("RGB", (720, 480), (9, 17, 31))
    from PIL import ImageDraw
    draw = ImageDraw.Draw(canvas)
    draw.rounded_rectangle((34, 34, 686, 446), radius=24, outline=(34, 211, 238), width=4)
    draw.text((70, 165), title[:28], fill=(235, 250, 255))
    draw.text((70, 225), subtitle[:58], fill=(129, 230, 217))
    out = io.BytesIO()
    canvas.save(out, format="JPEG", quality=88)
    return out.getvalue()


def _document_text(content: bytes, mime_type: str) -> str:
    stream = io.BytesIO(content)
    if mime_type == "application/pdf":
        return "\n".join((page.extract_text() or "") for page in PdfReader(stream).pages)
    if mime_type.endswith("wordprocessingml.document"):
        return "\n".join(p.text for p in Document(stream).paragraphs)
    if mime_type.endswith("presentationml.presentation"):
        deck = Presentation(stream)
        return "\n".join(shape.text for slide in deck.slides for shape in slide.shapes if hasattr(shape, "text"))
    if mime_type.endswith("spreadsheetml.sheet"):
        from openpyxl import load_workbook

        workbook = load_workbook(stream, read_only=True, data_only=True)
        lines: list[str] = []
        for sheet in workbook.worksheets:
            lines.append(f"[{sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                values = [str(value) for value in row if value not in (None, "")]
                if values:
                    lines.append("\t".join(values))
        workbook.close()
        return "\n".join(lines)
    return content.decode("utf-8-sig", errors="replace")


def inspect_media(content: bytes, mime_type: str, filename: str = "") -> ImageInspection:
    """Inspect an image, video or document while keeping its binary on disk."""
    if mime_type in IMAGE_MIME:
        return inspect_image(content, mime_type)
    if mime_type not in ALLOWED_MIME:
        raise VisualAssetError(
            "只支援 JPEG、PNG、WebP、GIF、AVIF、JFIF、MP4、MOV、WebM、PDF、DOCX、PPTX、XLSX 與文字文件",
            code="unsupported_media_type",
        )
    if not content:
        raise VisualAssetError("素材內容是空的", code="empty_file")
    digest = hashlib.sha256(content).hexdigest()
    if mime_type in DOCUMENT_MIME:
        try:
            text = _document_text(content, mime_type)[:20000]
        except Exception as exc:
            raise VisualAssetError("文件內容無法讀取", code="invalid_document") from exc
        return ImageInspection(
            width=0, height=0, orientation="document", exif_date="", sha256=digest,
            perceptual_hash=digest[:16], color_embedding=[], blur_score=0,
            brightness_score=0, quality_score=50 if text.strip() else 25,
            suitability={"poster": False, "video": False, "instagram_cover": False,
                         "reasons": ["文件素材", *([] if text.strip() else ["未擷取到內嵌文字"])]},
            thumbnail=_placeholder_thumbnail("DOCUMENT", filename or mime_type),
            asset_type="document", extracted_text=text,
        )
    return ImageInspection(
        width=0, height=0, orientation="video", exif_date="", sha256=digest,
        perceptual_hash=digest[:16], color_embedding=[], blur_score=0,
        brightness_score=0, quality_score=0,
        suitability={"poster": False, "video": True, "instagram_cover": False,
                     "reasons": ["影片已保留；影格解碼器未配置"]},
        thumbnail=_placeholder_thumbnail("VIDEO", filename or mime_type), asset_type="video",
    )


def semantic_embedding(text: str, dimensions: int = 192) -> list[float]:
    """不依賴外部服務的可重現語意特徵；AI 標籤/OCR 進來後會重算。"""
    normalized = re.sub(r"\s+", " ", text.lower()).strip()
    tokens = re.findall(r"[a-z0-9_-]{2,}|[\u3400-\u9fff]", normalized)
    tokens += [normalized[i:i + 2] for i in range(max(0, len(normalized) - 1)) if not normalized[i:i + 2].isspace()]
    vec = [0.0] * dimensions
    for token in tokens:
        digest = hashlib.blake2b(token.encode("utf-8"), digest_size=8).digest()
        slot = int.from_bytes(digest[:4], "big") % dimensions
        sign = 1.0 if digest[4] & 1 else -1.0
        vec[slot] += sign
    norm = math.sqrt(sum(v * v for v in vec)) or 1.0
    return [round(v / norm, 7) for v in vec]


def cosine(a: list[float], b: list[float]) -> float:
    if not a or not b or len(a) != len(b):
        return 0.0
    return sum(x * y for x, y in zip(a, b))


def hamming_hash(a: str, b: str) -> int:
    try:
        return (int(a, 16) ^ int(b, 16)).bit_count()
    except (TypeError, ValueError):
        return 64


class VisualAssetStore:
    def __init__(self, path: Path | None = None) -> None:
        self.path = Path(path or config.DB_PATH)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self._lock = threading.RLock()
        self._conn = sqlite3.connect(str(self.path), check_same_thread=False)
        self._conn.row_factory = sqlite3.Row
        with self._lock:
            apply_visual_migrations(self._conn, SCHEMA)
            stamp = now()
            self._conn.execute(
                """UPDATE visual_analysis_jobs SET status='failed',stage='interrupted',progress=100,
                       error_code='interrupted_process',error_message='服務重新啟動，請重試分析',updated_at=?
                   WHERE status='running'""",(stamp,),
            )
            # Do not overwrite review_status: a verified/probable asset must
            # not become failed just because the process restarted mid-job.
            self._conn.execute(
                """UPDATE visual_assets SET processing_state='failed',updated_at=?
                   WHERE id IN (SELECT asset_id FROM visual_analysis_jobs WHERE error_code='interrupted_process')
                     AND processing_state IN ('running','pending','queued','')""",(stamp,),
            )
            for name in SCENE_NAMES:
                self._conn.execute(
                    "INSERT OR IGNORE INTO visual_scenes(id,name,category,created_at) VALUES(?,?,?,?)",
                    ("scene_" + hashlib.sha1(name.encode()).hexdigest()[:12], name, name, now()),
                )
            self._conn.commit()

    def close(self) -> None:
        with self._lock:
            self._conn.close()

    def _rows(self, sql: str, params: tuple[Any, ...] = ()) -> list[dict[str, Any]]:
        with self._lock:
            rows = self._conn.execute(sql, params).fetchall()
        return [dict(row) for row in rows]

    def ensure_school(self, name: str) -> str | None:
        canonical = re.sub(r"\s+", "", (name or "").strip())
        if not canonical:
            return None
        aliases = {
            "淡江": "淡江大學", "淡大": "淡江大學", "TKU": "淡江大學",
            "政大": "國立政治大學", "政治大學": "國立政治大學", "NCCU": "國立政治大學",
        }
        canonical = aliases.get(canonical, canonical)
        with self._lock:
            row = self._conn.execute("SELECT id FROM visual_schools WHERE canonical_name=?", (canonical,)).fetchone()
            if row:
                return str(row[0])
            sid = new_id("school")
            self._conn.execute(
                "INSERT INTO visual_schools(id,canonical_name,aliases,created_at) VALUES(?,?,?,?)",
                (sid, canonical, "[]", now()),
            )
            self._conn.commit()
        return sid

    def ensure_club(self, school_id: str | None, name: str, *, source: dict[str, Any] | None = None, confirmed: bool = False) -> str | None:
        label = (name or "").strip()
        if not label:
            return None
        if not school_id:
            raise VisualAssetError("社團必須指定學校，避免不同學校同名社團混用", code="school_required")
        with self._lock:
            row = self._conn.execute("SELECT id FROM visual_clubs WHERE school_id=? AND name=?", (school_id, label)).fetchone()
            if row:
                return str(row[0])
            cid = new_id("club")
            stamp = now()
            self._conn.execute(
                "INSERT INTO visual_clubs(id,school_id,name,status,source,created_at,updated_at) VALUES(?,?,?,?,?,?,?)",
                (cid, school_id, label, "verified" if confirmed else "pending_review", dumps(source or {}), stamp, stamp),
            )
            self._conn.commit()
        return cid

    def create_asset(
        self, *, user_id: str, filename: str, mime_type: str, content: bytes,
        source: str = "upload", school: str = "", club: str = "", privacy: str = "private",
        commercial_use: str = "unknown", file_created_date: str = "", source_metadata: dict[str, Any] | None = None,
        relative_path: str = "", manifest_id: str = "", supersedes_asset_id: str = "",
        external_path: str = "", sha256_override: str = "", project_id: str = "",
    ) -> dict[str, Any]:
        external_source: Path | None = None
        if external_path:
            try:
                external_source = Path(external_path).resolve(strict=True)
            except (OSError, RuntimeError) as exc:
                raise VisualAssetError("外部素材來源不存在或無法讀取", code="source_unavailable") from exc
            if not external_source.is_file():
                raise VisualAssetError("外部素材來源不是檔案", code="source_unavailable")
        if len(content) > config.VISUAL_MAX_FILE_BYTES and not external_source:
            raise VisualAssetError(f"素材超過 {config.VISUAL_MAX_FILE_BYTES // 1_000_000}MB 上限", code="file_too_large")
        if privacy not in {"private", "shared", "public"}:
            raise VisualAssetError("隱私狀態必須是 private、shared 或 public")
        if commercial_use not in {"allowed", "not_allowed", "unknown"}:
            raise VisualAssetError("商用權限必須是 allowed、not_allowed 或 unknown")
        inspection = inspect_media(content, mime_type, filename)
        school_id = self.ensure_school(school)
        club_id = self.ensure_club(school_id, club, source={"type": "user_input"}, confirmed=True) if club else None
        asset_id = new_id("asset")
        safe_ext = ALLOWED_MIME[mime_type]
        target_dir = (config.VISUAL_ASSET_DIR / user_id / asset_id).resolve()
        root = config.VISUAL_ASSET_DIR.resolve()
        if root not in target_dir.parents:
            raise VisualAssetError("資產儲存路徑無效")
        target_dir.mkdir(parents=True, exist_ok=False)
        original_path = external_source or (target_dir / ("original" + safe_ext))
        thumb_path = target_dir / "thumbnail.jpg"
        # 外部匯入只寫入素材庫縮圖；原始檔維持在原位置，避免 9GB 級資料
        # 再複製一份，也不會有任何 API 路徑覆寫或刪除原檔。
        if not external_source:
            original_path.write_bytes(content)
        thumb_path.write_bytes(inspection.thumbnail)
        digest = (sha256_override or inspection.sha256).strip()
        metadata = dict(source_metadata or {})
        if external_source:
            metadata.setdefault("storage_mode", "external")
        duplicate = self._rows(
            "SELECT id FROM visual_assets WHERE sha256=? AND (user_id=? OR privacy IN ('shared','public')) ORDER BY created_at LIMIT 1",
            (digest, user_id),
        )
        duplicate_of = duplicate[0]["id"] if duplicate else None
        stamp = now()
        with self._lock:
            self._conn.execute(
                """INSERT INTO visual_assets(
                    id,user_id,original_filename,storage_path,thumbnail_path,mime_type,extension,width,height,orientation,
                    uploaded_at,source,source_metadata,exif_date,file_created_date,sha256,perceptual_hash,color_embedding,
                    semantic_embedding,privacy,commercial_use,analysis_status,quality_score,blur_score,brightness_score,
                    suitability,ocr_text,review_status,school_id,club_id,duplicate_of,created_at,updated_at,
                    asset_type,relative_path,manifest_id,aspect_ratio,people_count,overall_confidence,processing_state,
                    last_confirmed_at,supersedes_asset_id,project_id,owner_id
                ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    asset_id,user_id,filename[:240],str(original_path),str(thumb_path),mime_type,safe_ext,
                    inspection.width,inspection.height,inspection.orientation,stamp,source[:120],dumps(metadata),
                    inspection.exif_date,file_created_date[:40],digest,inspection.perceptual_hash,
                    dumps(inspection.color_embedding),dumps(semantic_embedding(filename + " " + inspection.extracted_text)),privacy,commercial_use,
                    "local_complete",inspection.quality_score,inspection.blur_score,inspection.brightness_score,
                    dumps(inspection.suitability),inspection.extracted_text,"pending_review",school_id,club_id,duplicate_of,stamp,stamp,
                    inspection.asset_type,relative_path[:1000],manifest_id[:100],
                    round(inspection.width / inspection.height, 6) if inspection.height else 0,
                    0,0,"completed","",supersedes_asset_id[:80],(project_id or "")[:160],user_id,
                ),
            )
            for value, date_source, confidence in (
                (inspection.exif_date, "exif", 0.96), (file_created_date, "file_metadata", 0.65),
            ):
                normalized = normalize_date(value)
                if normalized:
                    self._conn.execute(
                        "INSERT INTO visual_date_candidates(id,asset_id,value,source,confidence,status,evidence,created_at) VALUES(?,?,?,?,?,?,?,?)",
                        (new_id("date"),asset_id,normalized,date_source,confidence,"probable",value,stamp),
                    )
            self._conn.commit()
        return self.get_asset(asset_id, user_id) or {}

    def visible_asset(self, asset_id: str, user_id: str) -> dict[str, Any] | None:
        rows = self._rows(
            """SELECT a.*,COALESCE(s.canonical_name,'') AS school_name,COALESCE(c.name,'') AS club_name
               FROM visual_assets a LEFT JOIN visual_schools s ON s.id=a.school_id LEFT JOIN visual_clubs c ON c.id=a.club_id
               WHERE a.id=? AND (a.user_id=? OR a.privacy IN ('shared','public'))""",
            (asset_id, user_id),
        )
        return rows[0] if rows else None

    def owned_asset(self, asset_id: str, user_id: str) -> dict[str, Any] | None:
        rows = self._rows(
            """SELECT a.*,COALESCE(s.canonical_name,'') AS school_name,COALESCE(c.name,'') AS club_name
               FROM visual_assets a LEFT JOIN visual_schools s ON s.id=a.school_id LEFT JOIN visual_clubs c ON c.id=a.club_id
               WHERE a.id=? AND a.user_id=?""",(asset_id,user_id),
        )
        return rows[0] if rows else None

    def get_asset(self, asset_id: str, user_id: str) -> dict[str, Any] | None:
        asset = self.visible_asset(asset_id, user_id)
        if not asset:
            return None
        observations = self._rows(
            "SELECT * FROM visual_observations WHERE asset_id=? ORDER BY entity_type,status,confidence DESC", (asset_id,),
        )
        dates = self._rows(
            "SELECT * FROM visual_date_candidates WHERE asset_id=? ORDER BY value,confidence DESC", (asset_id,),
        )
        jobs = self._rows(
            "SELECT id,status,stage,progress,error_code,error_message,updated_at FROM visual_analysis_jobs WHERE asset_id=? ORDER BY created_at DESC LIMIT 1",
            (asset_id,),
        )
        return self.public_asset(asset, observations=observations, dates=dates, job=(jobs[0] if jobs else None))

    def public_asset(
        self, asset: dict[str, Any], *, observations: list[dict[str, Any]] | None = None,
        dates: list[dict[str, Any]] | None = None, job: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        result = {k: asset.get(k) for k in (
            "id","original_filename","mime_type","width","height","orientation","uploaded_at","source",
            "exif_date","file_created_date","privacy","commercial_use","analysis_status","quality_score",
            "blur_score","brightness_score","ocr_text","review_status","duplicate_of","school_id","club_id","created_at","updated_at",
            "asset_type","relative_path","manifest_id","aspect_ratio","people_count","overall_confidence",
            "processing_state","last_confirmed_at","supersedes_asset_id","project_id",
        )}
        result["asset_id"] = result.pop("id")
        result["thumbnail_url"] = f"/api/visual-assets/{result['asset_id']}/file?variant=thumbnail"
        result["original_url"] = f"/api/visual-assets/{result['asset_id']}/file?variant=original"
        result["suitability"] = loads(asset.get("suitability"), {})
        result["analysis"] = loads(asset.get("analysis"), {})
        result["source_metadata"] = loads(asset.get("source_metadata"), {})
        result["school_name"] = asset.get("school_name","")
        result["club_name"] = asset.get("club_name","")
        if observations is not None:
            for obs in observations:
                obs["evidence"] = loads(obs.get("evidence"), {})
            result["observations"] = observations
        if dates is not None:
            distinct = sorted({d["value"] for d in dates if d.get("status") != "failed"})
            result["date_candidates"] = dates
            result["date_status"] = "conflicted" if len(distinct) > 1 else ("verified" if any(d.get("status") == "verified" for d in dates) else "probable" if dates else "pending_review")
        if job is not None:
            result["analysis_job"] = job
        return result

    def _is_external_asset(self, asset: dict[str, Any]) -> bool:
        metadata = loads(asset.get("source_metadata"), {})
        return isinstance(metadata, dict) and metadata.get("storage_mode") == "external"

    def _confine_visual_file(
        self, raw: Path, mime: str, filename: str, *, allow_external: bool,
    ) -> tuple[Path, str, str] | None:
        """Serve only files we manage, or original bytes of an external import.

        Local assets must stay under VISUAL_ASSET_DIR. External originals may
        live outside that tree; derivatives still cannot. Missing files return
        None (HTTP 404) rather than leaking path validity.
        """
        try:
            resolved = Path(raw).resolve()
        except (OSError, RuntimeError):
            return None
        try:
            root = config.VISUAL_ASSET_DIR.resolve()
        except (OSError, RuntimeError):
            return None
        under_library = resolved == root or root in resolved.parents
        if not under_library:
            if not allow_external or not resolved.is_file():
                return None
            try:
                db_path = Path(config.DB_PATH).resolve()
            except (OSError, RuntimeError):
                db_path = None
            if db_path is not None and resolved == db_path:
                return None
            app_root = (config.ROOT / "app").resolve()
            git_root = (config.ROOT / ".git").resolve()
            if app_root in resolved.parents or git_root in resolved.parents or resolved == app_root:
                return None
            return resolved, mime, filename
        if not resolved.is_file():
            return None
        return resolved, mime, filename

    def asset_file(self, asset_id: str, user_id: str, variant: str) -> tuple[Path, str, str] | None:
        asset = self.visible_asset(asset_id,user_id)
        if not asset:
            return None
        external = self._is_external_asset(asset)
        if variant == "original":
            return self._confine_visual_file(
                Path(asset["storage_path"]), asset["mime_type"], asset["original_filename"], allow_external=external,
            )
        if variant == "thumbnail":
            return self._confine_visual_file(
                Path(asset["thumbnail_path"]), "image/jpeg", f"{asset_id}-thumbnail.jpg", allow_external=False,
            )
        ratios = {"1:1":(1,1),"4:5":(4,5),"9:16":(9,16),"16:9":(16,9)}
        if variant not in ratios:
            raise VisualAssetError("variant 必須是 original、thumbnail、1:1、4:5、9:16 或 16:9")
        if asset.get("asset_type") != "image":
            raise VisualAssetError("文件與影片目前只支援原始檔及縮圖輸出", code="unsupported_rendition")
        confined = self._confine_visual_file(
            Path(asset["storage_path"]), asset["mime_type"], asset["original_filename"], allow_external=external,
        )
        if not confined:
            return None
        source = confined[0]
        metadata = loads(asset.get("source_metadata"), {})
        # External imports must never write a rendition beside the user's
        # original file.  Keep all generated derivatives under our managed
        # visual-assets directory instead.
        target_parent = Path(asset["thumbnail_path"]).parent if isinstance(metadata, dict) and metadata.get("storage_mode") == "external" else source.parent
        try:
            target_parent = target_parent.resolve()
        except (OSError, RuntimeError):
            return None
        library_root = config.VISUAL_ASSET_DIR.resolve()
        if library_root not in target_parent.parents and target_parent != library_root:
            return None
        target = target_parent / f"rendition-{variant.replace(':','x')}.jpg"
        if not target.exists():
            with Image.open(source) as raw:
                image = ImageOps.exif_transpose(raw).convert("RGB")
                target_ratio = ratios[variant][0] / ratios[variant][1]
                current_ratio = image.width / image.height
                if current_ratio > target_ratio:
                    wanted = round(image.height * target_ratio)
                    left = max(0,(image.width - wanted)//2)
                    image = image.crop((left,0,left+wanted,image.height))
                else:
                    wanted = round(image.width / target_ratio)
                    top = max(0,(image.height - wanted)//2)
                    image = image.crop((0,top,image.width,top+wanted))
                max_width = 1920 if variant == "16:9" else 1440
                if image.width > max_width:
                    image.thumbnail((max_width,round(max_width/target_ratio)),Image.Resampling.LANCZOS)
                # 衍生圖可重建；原圖永不被開啟為寫入目標。
                image.save(target,format="JPEG",quality=91,optimize=True)
        return self._confine_visual_file(
            target, "image/jpeg", f"{Path(asset['original_filename']).stem}-{variant.replace(':','x')}.jpg",
            allow_external=False,
        )

    def export_package(self, user_id: str, asset_ids: list[str], *, rendition: str = "original") -> Path:
        unique_ids = list(dict.fromkeys(asset_ids))[:100]
        if not unique_ids:
            raise VisualAssetError("至少選擇一張圖片")
        records: list[dict[str, Any]] = []
        files: list[tuple[Path,str]] = []
        for asset_id in unique_ids:
            item = self.get_asset(asset_id,user_id)
            resolved = self.asset_file(asset_id,user_id,rendition)
            if not item or not resolved:
                continue
            path,_,filename = resolved
            records.append(item)
            files.append((path,filename))
        if not records:
            raise VisualAssetError("選取的圖片不存在或沒有權限",code="not_found")
        config.VISUAL_EXPORT_DIR.mkdir(parents=True,exist_ok=True)
        out = config.VISUAL_EXPORT_DIR / f"visual-assets-{datetime.now().strftime('%Y%m%d-%H%M%S')}-{secrets.token_hex(3)}.zip"
        csv_buffer = io.StringIO()
        writer = csv.writer(csv_buffer)
        writer.writerow(["asset_id","filename","width","height","uploaded_at","school_id","club_id","quality_score","privacy","commercial_use","ocr_text"])
        for item in records:
            writer.writerow([item.get(k,"") for k in ("asset_id","original_filename","width","height","uploaded_at","school_id","club_id","quality_score","privacy","commercial_use","ocr_text")])
        with zipfile.ZipFile(out,"w",compression=zipfile.ZIP_DEFLATED) as archive:
            for path,filename in files:
                archive.write(path,arcname=f"images/{filename}")
            archive.writestr("manifest.json",json.dumps(records,ensure_ascii=False,indent=2))
            archive.writestr("manifest.csv",csv_buffer.getvalue().encode("utf-8-sig"))
        return out

    def create_job(self, asset_id: str) -> str:
        jid = new_id("job")
        stamp = now()
        with self._lock:
            attempts = int(self._conn.execute("SELECT COUNT(*) FROM visual_analysis_jobs WHERE asset_id=?",(asset_id,)).fetchone()[0]) + 1
            self._conn.execute(
                "INSERT INTO visual_analysis_jobs(id,asset_id,status,stage,progress,attempts,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?)",
                (jid,asset_id,"running","preparing",5,attempts,stamp,stamp),
            )
            self._conn.execute("UPDATE visual_assets SET analysis_status='analyzing',processing_state='running',updated_at=? WHERE id=?", (stamp,asset_id))
            self._conn.commit()
        return jid

    def update_job(self, job_id: str, *, status: str, stage: str, progress: int, error_code: str = "", error_message: str = "") -> None:
        with self._lock:
            self._conn.execute(
                "UPDATE visual_analysis_jobs SET status=?,stage=?,progress=?,error_code=?,error_message=?,updated_at=? WHERE id=?",
                (status,stage,max(0,min(100,progress)),error_code[:80],error_message[:500],now(),job_id),
            )
            if status == "failed":
                self._conn.execute(
                    "UPDATE visual_assets SET processing_state='failed',review_status='failed',updated_at=? WHERE id=(SELECT asset_id FROM visual_analysis_jobs WHERE id=?)",
                    (now(),job_id),
                )
            elif status == "complete":
                self._conn.execute(
                    "UPDATE visual_assets SET processing_state='completed',updated_at=? WHERE id=(SELECT asset_id FROM visual_analysis_jobs WHERE id=?)",
                    (now(),job_id),
                )
            self._conn.commit()

    def set_analysis(self, asset_id: str, result: dict[str, Any], *, status: str = "complete") -> None:
        summary = " ".join([
            str(result.get("summary") or ""), str(result.get("ocr_text") or ""),
            " ".join(str(v) for v in result.get("objects", []) if v),
            " ".join(str((v or {}).get("label") or "") for v in result.get("scenes", []) if isinstance(v, dict)),
        ])
        with self._lock:
            self._conn.execute(
                "UPDATE visual_assets SET ocr_text=?,analysis=?,semantic_embedding=?,analysis_status=?,people_count=?,overall_confidence=?,review_status=CASE WHEN review_status IN ('pending_review','failed') THEN 'probable' ELSE review_status END,updated_at=? WHERE id=?",
                (str(result.get("ocr_text") or "")[:20000],dumps(result),dumps(semantic_embedding(summary)),status,
                 int((result.get("people") or {}).get("count") or 0) if isinstance(result.get("people"),dict) else 0,
                 max([float(x.get("confidence") or 0) for key in ("scenes","clubs") for x in (result.get(key) or []) if isinstance(x,dict)] or [0.0]),
                 now(),asset_id),
            )
            self._conn.commit()

    def mark_analysis_status(self, asset_id: str, status: str) -> None:
        with self._lock:
            self._conn.execute("UPDATE visual_assets SET analysis_status=?,updated_at=? WHERE id=?",(status,now(),asset_id))
            self._conn.commit()

    def add_observation(
        self, asset_id: str, entity_type: str, *, label: str = "", entity_id: str = "",
        status: str = "probable", confidence: float = 0.0, source: str, evidence: dict[str, Any] | None = None,
    ) -> str:
        if status not in STATUS_VALUES:
            status = "pending_review"
        oid = new_id("obs")
        stamp = now()
        with self._lock:
            # 同一輪重跑分析時不累積完全相同的 AI observation。
            current = self._conn.execute(
                "SELECT id FROM visual_observations WHERE asset_id=? AND entity_type=? AND entity_id=? AND label=? AND source=?",
                (asset_id,entity_type,entity_id,label,source),
            ).fetchone()
            if current:
                self._conn.execute(
                    "UPDATE visual_observations SET status=?,confidence=?,evidence=?,updated_at=? WHERE id=?",
                    (status,max(0,min(1,float(confidence))),dumps(evidence or {}),stamp,current[0]),
                )
                oid = str(current[0])
            else:
                self._conn.execute(
                    "INSERT INTO visual_observations(id,asset_id,entity_type,entity_id,label,status,confidence,source,evidence,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)",
                    (oid,asset_id,entity_type,entity_id,label[:300],status,max(0,min(1,float(confidence))),source[:120],dumps(evidence or {}),stamp,stamp),
                )
            self._conn.commit()
        return oid

    def add_date(self, asset_id: str, value: str, *, source: str, confidence: float, evidence: str = "", status: str = "probable") -> str | None:
        normalized = normalize_date(value)
        if not normalized:
            return None
        with self._lock:
            exists = self._conn.execute(
                "SELECT id FROM visual_date_candidates WHERE asset_id=? AND value=? AND source=?", (asset_id,normalized,source),
            ).fetchone()
            if exists:
                return str(exists[0])
            did = new_id("date")
            self._conn.execute(
                "INSERT INTO visual_date_candidates(id,asset_id,value,source,confidence,status,evidence,created_at) VALUES(?,?,?,?,?,?,?,?)",
                (did,asset_id,normalized,source,max(0,min(1,float(confidence))),status,evidence[:1000],now()),
            )
            distinct = int(self._conn.execute(
                "SELECT COUNT(DISTINCT value) FROM visual_date_candidates WHERE asset_id=? AND status!='failed'",(asset_id,),
            ).fetchone()[0])
            if distinct > 1:
                self._conn.execute("UPDATE visual_assets SET review_status='conflicted',updated_at=? WHERE id=?",(now(),asset_id))
            self._conn.commit()
        return did

    def scene_id(self, label: str) -> str:
        name = (label or "其他場景").strip()[:80]
        with self._lock:
            row = self._conn.execute("SELECT id FROM visual_scenes WHERE name=?", (name,)).fetchone()
            if row:
                return str(row[0])
            sid = new_id("scene")
            self._conn.execute(
                "INSERT INTO visual_scenes(id,name,category,created_at) VALUES(?,?,?,?)",
                (sid,name,"custom",now()),
            )
            self._conn.commit()
        return sid

    def ensure_event(
        self, *, school_id: str | None, club_id: str | None, name: str, event_type: str = "",
        event_date: str = "", start_time: str = "", end_time: str = "", location: str = "",
        status: str = "probable", evidence: list[dict[str, Any]] | None = None,
    ) -> str:
        label = (name or "待確認活動").strip()[:200]
        date_value = normalize_date(event_date)
        with self._lock:
            # Re-analysis of the same source asset is idempotent. A matching
            # label/date from another asset is still never permission to merge.
            source_asset = next((str(item.get("asset_id") or "") for item in (evidence or []) if isinstance(item,dict) and item.get("asset_id")),"")
            if source_asset:
                row = self._conn.execute(
                    "SELECT id FROM visual_events WHERE evidence LIKE ? ORDER BY created_at LIMIT 1",(f'%"asset_id":"{source_asset}"%',),
                ).fetchone()
                if row:
                    return str(row[0])
            eid = new_id("event")
            stamp = now()
            self._conn.execute(
                """INSERT INTO visual_events(id,school_id,club_id,name,event_type,event_date,start_time,end_time,location,status,evidence,created_at,updated_at)
                   VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (eid,school_id,club_id,label,event_type[:100],date_value,normalize_time(start_time),normalize_time(end_time),location[:300],status,dumps(evidence or []),stamp,stamp),
            )
            self._conn.commit()
        return eid

    def confirmed_person_references(self, user_id: str, *, limit: int = 8) -> list[dict[str, Any]]:
        return self._rows(
            """SELECT p.id AS person_id,p.name,a.storage_path,a.mime_type
               FROM visual_person_references r
               JOIN visual_people p ON p.id=r.person_id AND p.status='verified'
               JOIN visual_assets a ON a.id=r.asset_id
               WHERE r.status='verified' AND (a.user_id=? OR a.privacy IN ('shared','public'))
               ORDER BY r.created_at DESC LIMIT ?""",
            (user_id,limit),
        )

    def search(
        self, user_id: str, *, query: str = "", page: int = 1, limit: int = 30,
        school: str = "", club: str = "", person: str = "", scene: str = "", event: str = "",
        date_from: str = "", date_to: str = "", ratio: str = "", quality_min: float = 0,
        commercial_use: str = "", privacy: str = "", duplicate: str = "",
        people_min: int = 0, brightness_min: float = 0, verification_status: str = "",
    ) -> tuple[list[dict[str, Any]], int, dict[str, Any]]:
        q = query.strip().lower()
        school_filter = school or infer_school(query)
        # Natural-language searches should use the same conservative labels
        # produced by the data organizer (for example ``領袖禪學社`` from a
        # curated 淡大劇本 folder), while explicit filters still win.
        club_filter = club or infer_club(query)
        # Generated thumbnails/renditions are retained for lineage but are not
        # independent search results.  They were never user-authored assets.
        where = ["COALESCE(a.source,'')!='derived_thumbnail'", "(a.user_id=? OR a.privacy IN ('shared','public'))"]
        params: list[Any] = [user_id]
        if school_filter:
            where.append("(a.school_id IN (SELECT id FROM visual_schools WHERE canonical_name LIKE ? OR aliases LIKE ?) OR a.id IN (SELECT asset_id FROM visual_observations WHERE entity_type='school' AND label LIKE ? AND review_action!='ignored'))")
            params.extend([f"%{school_filter}%", f"%{school_filter}%", f"%{school_filter}%"])
        if club_filter:
            where.append("a.id IN (SELECT asset_id FROM visual_observations WHERE entity_type='club' AND label LIKE ? AND review_action!='ignored') OR a.club_id IN (SELECT id FROM visual_clubs WHERE name LIKE ? OR aliases LIKE ?)")
            params.extend([f"%{club_filter}%",f"%{club_filter}%",f"%{club_filter}%"])
        if person:
            # A real name is a hard identity filter: only user-confirmed
            # observations tied to a verified person are eligible.
            where.append("a.id IN (SELECT o.asset_id FROM visual_observations o JOIN visual_people p ON p.id=o.entity_id AND p.status='verified' WHERE o.entity_type='person' AND o.status='verified' AND o.review_action!='ignored' AND (p.name LIKE ? OR p.nickname LIKE ? OR o.label LIKE ?))")
            params.extend([f"%{person}%",f"%{person}%",f"%{person}%"])
        if scene:
            where.append("a.id IN (SELECT asset_id FROM visual_observations WHERE entity_type='scene' AND label LIKE ? AND review_action!='ignored')")
            params.append(f"%{scene}%")
        if event:
            where.append("a.id IN (SELECT o.asset_id FROM visual_observations o LEFT JOIN visual_events e ON e.id=o.entity_id WHERE o.entity_type='event' AND o.review_action!='ignored' AND (e.name LIKE ? OR o.label LIKE ?))")
            params.extend([f"%{event}%",f"%{event}%"])
        if date_from:
            where.append("EXISTS (SELECT 1 FROM visual_date_candidates d WHERE d.asset_id=a.id AND d.status!='failed' AND d.value>=?)")
            params.append(date_from)
        if date_to:
            where.append("EXISTS (SELECT 1 FROM visual_date_candidates d WHERE d.asset_id=a.id AND d.status!='failed' AND d.value<=?)")
            params.append(date_to)
        if quality_min:
            where.append("a.quality_score>=?")
            params.append(float(quality_min))
        inferred_people_min = max(int(people_min or 0), 2 if any(term in q for term in ("多人", "合照", "團體照")) else 0)
        inferred_quality_min = max(float(quality_min or 0), 70.0 if any(term in q for term in ("高清", "高畫質", "畫質最好", "人物清楚", "清晰")) else 0.0)
        if inferred_quality_min > float(quality_min or 0):
            where.append("a.quality_score>=?")
            params.append(inferred_quality_min)
        if inferred_people_min:
            where.append("a.people_count>=?")
            params.append(inferred_people_min)
        if brightness_min:
            where.append("a.brightness_score>=?")
            params.append(float(brightness_min))
        if verification_status:
            where.append("a.review_status=?")
            params.append(verification_status)
        if commercial_use:
            where.append("a.commercial_use=?")
            params.append(commercial_use)
        if privacy:
            where.append("a.privacy=?")
            params.append(privacy)
        if duplicate == "only":
            where.append("a.duplicate_of IS NOT NULL")
        elif duplicate == "exclude":
            where.append("a.duplicate_of IS NULL")
        condition = " AND ".join(f"({item})" for item in where)
        rows = self._rows(
            f"""SELECT a.*,
                   COALESCE(s.canonical_name,'') AS school_name,
                   COALESCE(c.name,'') AS club_name,
                   COALESCE((SELECT GROUP_CONCAT(label,' ') FROM visual_observations o WHERE o.asset_id=a.id AND o.review_action!='ignored'),'') AS labels,
                   COALESCE((SELECT GROUP_CONCAT(value,' ') FROM visual_date_candidates d WHERE d.asset_id=a.id AND d.status!='failed'),'') AS dates,
                   COALESCE((SELECT SUM(CASE WHEN event_type IN ('selected','used','downloaded') THEN 1 WHEN event_type='excluded' THEN -1 ELSE 0 END) FROM visual_learning_events l WHERE l.asset_id=a.id),0) AS feedback
               FROM visual_assets a
               LEFT JOIN visual_schools s ON s.id=a.school_id
               LEFT JOIN visual_clubs c ON c.id=a.club_id
               WHERE {condition} ORDER BY a.uploaded_at DESC LIMIT 600""",
            tuple(params),
        )
        qvec = semantic_embedding(q) if q else []
        qtokens = [t for t in re.split(r"[\s,，、。]+", q) if len(t) >= 2]
        compact_q = re.sub(r"找|照片|圖片|素材|請|幫我|的", "", q)
        qtokens += [compact_q[i:i+2] for i in range(max(0,len(compact_q)-1)) if not compact_q[i:i+2].isspace()]
        qtokens = list(dict.fromkeys(qtokens))
        ratio_alias = infer_ratio(q, ratio)
        # Keep the original hard-scene contract for explicit visual scenes;
        # directory-derived event names such as 茶會／社課 remain searchable
        # through labels and paths without excluding legacy filename matches.
        hard_scene_names = ("活動現場","報到區","校園","教室","禪堂","舞台","戶外","海報","文件")
        scene_intent_names = (*hard_scene_names, "茶會", "講座", "社課", "聚餐")
        inferred_scene = scene or next((name for name in scene_intent_names if name in q), "")
        hard_scene = scene or next((name for name in hard_scene_names if name in q), "")
        candidates: list[tuple[dict[str, Any], int, float, list[str]]] = []
        for row in rows:
            if ratio_alias and not ratio_matches(row["width"], row["height"], ratio_alias):
                continue
            # 自然語言中明確出現既知場景時，把它當證據條件，不把只有學校名稱
            # 相符、卻沒有任何場景觀察的圖片冒充成「校園照」。
            if hard_scene and hard_scene not in str(row.get("labels") or ""):
                continue
            document = " ".join(str(row.get(key) or "") for key in ("original_filename","relative_path","ocr_text","school_name","club_name","labels","dates"))
            doc_lower = document.lower()
            keyword_hits = sum(1 for token in qtokens if token in doc_lower)
            sem = cosine(qvec, loads(row.get("semantic_embedding"), [])) if q else 0.0
            reasons: list[str] = []
            if keyword_hits:
                reasons.append(f"符合 {keyword_hits} 個文字／實體條件")
            if sem > 0.12:
                reasons.append("圖像描述與查詢語意相近")
            if inferred_scene and inferred_scene in row.get("labels", ""):
                reasons.append(f"場景符合「{inferred_scene}」")
            if float(row.get("quality_score") or 0) >= 75:
                reasons.append("畫質評分高")
            if ratio_alias:
                reasons.append(f"比例符合 {ratio_alias}")
            if row.get("duplicate_of"):
                reasons.append("偵測到重複原圖")
            if q and keyword_hits == 0 and sem <= 0.02:
                continue
            candidates.append((row,keyword_hits,sem,reasons or ["符合目前篩選條件"]))
        keyword_ranking = [item[0]["id"] for item in sorted(candidates,key=lambda item:(item[1],item[0].get("uploaded_at","")),reverse=True) if item[1] > 0]
        semantic_ranking = [item[0]["id"] for item in sorted(candidates,key=lambda item:item[2],reverse=True) if item[2] > 0.02]
        fused = reciprocal_rank_fusion([keyword_ranking,semantic_ranking]) if q else {}
        ranked: list[tuple[float,dict[str,Any],list[str]]] = []
        for row,keyword_hits,sem,reasons in candidates:
            score = fused.get(row["id"],0.0) * 100.0
            score += float(row.get("quality_score") or 0) / 250.0
            score += max(-0.5,min(0.75,float(row.get("feedback") or 0) * 0.08))
            if inferred_scene:
                score += 1.0
            ranked.append((score,row,reasons))
        ranked.sort(key=lambda item: (item[0], item[1].get("uploaded_at", "")), reverse=True)
        total = len(ranked)
        chunk = ranked[(page - 1) * limit:page * limit]
        items = []
        for score,row,reasons in chunk:
            item = self.public_asset(row)
            item["match_score"] = round(score, 4)
            item["recommendation_reasons"] = reasons
            item["school_name"] = row.get("school_name", "")
            item["club_name"] = row.get("club_name", "")
            observations = self._rows(
                "SELECT entity_type,label,entity_id,status,confidence,source,evidence FROM visual_observations WHERE asset_id=? AND review_action!='ignored' ORDER BY status='verified' DESC,confidence DESC",
                (row["id"],),
            )
            for observation in observations:
                observation["evidence"] = loads(observation.get("evidence"), {})
            dates = self._rows(
                "SELECT value,source,confidence,status,evidence FROM visual_date_candidates WHERE asset_id=? AND status!='failed' ORDER BY status='verified' DESC,confidence DESC",
                (row["id"],),
            )
            item["matched_entities"] = observations
            item["matched_dates"] = dates
            item["confidence_score"] = round(float(row.get("overall_confidence") or 0), 4)
            item["evidence"] = {
                "asset_id": row["id"], "source": row.get("source") or "",
                "detail_url": f"/api/visual-assets/{row['id']}",
            }
            items.append(item)
        parsed = {
            "query": query, "school": school_filter, "club": club_filter, "person": person,
            "scene": inferred_scene, "event": event, "date_from": date_from, "date_to": date_to,
            "ratio": ratio_alias, "quality_min": inferred_quality_min, "commercial_use": commercial_use,
            "people_min": inferred_people_min, "brightness_min": brightness_min,
            "verification_status": verification_status,
        }
        return items,total,parsed

    def search_by_image(self, user_id: str, content: bytes, mime_type: str, *, page: int = 1, limit: int = 30) -> tuple[list[dict[str, Any]], int]:
        probe = inspect_image(content, mime_type)
        rows = self._rows(
            "SELECT * FROM visual_assets WHERE user_id=? OR privacy IN ('shared','public') ORDER BY uploaded_at DESC LIMIT 1000",
            (user_id,),
        )
        ranked = []
        for row in rows:
            distance = hamming_hash(probe.perceptual_hash, str(row.get("perceptual_hash") or ""))
            color = cosine(probe.color_embedding, loads(row.get("color_embedding"), []))
            similarity = max(0.0, min(1.0, (1.0 - distance / 64.0) * 0.78 + max(0.0,color) * 0.22))
            if similarity >= 0.35:
                ranked.append((similarity,row,distance))
        ranked.sort(key=lambda item: item[0], reverse=True)
        total = len(ranked)
        items = []
        for similarity,row,distance in ranked[(page-1)*limit:page*limit]:
            item = self.public_asset(row)
            item["match_score"] = round(similarity,4)
            item["recommendation_reasons"] = ["視覺構圖與色彩相近", f"感知雜湊距離 {distance}"]
            items.append(item)
        return items,total

    @staticmethod
    def safe_relative_path(value: str, filename: str) -> str:
        raw = (value or filename).replace("\\", "/").lstrip("/")
        path = PurePosixPath(raw)
        if not raw or ".." in path.parts:
            raise VisualAssetError("manifest relative_path 不可跳出匯入根目錄", code="invalid_manifest_path")
        return str(path)[:1000]

    def start_import(
        self, user_id: str, *, idempotency_key: str, root_name: str,
        manifest: dict[str, Any], total_count: int,
    ) -> dict[str, Any]:
        key = idempotency_key.strip()[:160]
        if len(key) < 8:
            raise VisualAssetError("idempotency_key 至少需要 8 個字元", code="invalid_idempotency_key")
        with self._lock:
            row = self._conn.execute(
                "SELECT id FROM visual_imports WHERE user_id=? AND idempotency_key=?", (user_id,key),
            ).fetchone()
            stamp = now()
            if row:
                import_id = str(row[0])
                self._conn.execute(
                    "UPDATE visual_imports SET manifest_json=?,total_count=MAX(total_count,?),updated_at=? WHERE id=?",
                    (dumps(manifest),total_count,stamp,import_id),
                )
            else:
                import_id = new_id("import")
                self._conn.execute(
                    """INSERT INTO visual_imports(id,user_id,idempotency_key,root_name,manifest_json,status,total_count,created_at,updated_at)
                       VALUES(?,?,?,?,?,'pending_review',?,?,?)""",
                    (import_id,user_id,key,root_name[:240],dumps(manifest),total_count,stamp,stamp),
                )
            self._conn.commit()
        return self.get_import(import_id,user_id) or {}

    def ingest_import_item(
        self, *, user_id: str, import_id: str, client_key: str, relative_path: str,
        filename: str, mime_type: str, content: bytes, last_modified: str = "",
        school: str = "", club: str = "", source: str = "folder_import",
        privacy: str = "private", commercial_use: str = "unknown", resync: bool = False,
        project_id: str = "",
    ) -> tuple[dict[str, Any], bool]:
        current_import = self._rows("SELECT * FROM visual_imports WHERE id=? AND user_id=?", (import_id,user_id))
        if not current_import:
            raise VisualAssetError("找不到匯入工作或沒有權限", code="not_found")
        rel = self.safe_relative_path(relative_path, filename)
        key = (client_key or rel)[:240]
        digest = hashlib.sha256(content).hexdigest()
        existing_rows = self._rows(
            "SELECT * FROM visual_import_items WHERE import_id=? AND client_key=?", (import_id,key),
        )
        existing = existing_rows[0] if existing_rows else None
        if existing and existing.get("status") == "verified" and existing.get("sha256") == digest:
            item = self.get_asset(str(existing.get("asset_id") or ""),user_id)
            if item:
                return item, True
        if existing and existing.get("asset_id") and existing.get("sha256") != digest and not resync:
            raise VisualAssetError("檔案內容已變更；請使用 resync=true 建立新版本", code="resync_required")
        stamp = now()
        item_id = str(existing.get("id")) if existing else new_id("import_item")
        attempts = int(existing.get("attempts") or 0) + 1 if existing else 1
        with self._lock:
            self._conn.execute(
                """INSERT INTO visual_import_items(
                       id,import_id,client_key,relative_path,filename,size_bytes,last_modified,sha256,status,attempts,created_at,updated_at
                   ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)
                   ON CONFLICT(import_id,client_key) DO UPDATE SET
                       relative_path=excluded.relative_path,filename=excluded.filename,size_bytes=excluded.size_bytes,
                       last_modified=excluded.last_modified,sha256=excluded.sha256,status='pending_review',
                       error_code='',error_message='',attempts=?,updated_at=?""",
                (item_id,import_id,key,rel,filename[:240],len(content),last_modified[:80],digest,"pending_review",attempts,stamp,stamp,attempts,stamp),
            )
            self._conn.commit()
        try:
            asset = self.create_asset(
                user_id=user_id,filename=filename,mime_type=mime_type,content=content,source=source,
                school=school,club=club,privacy=privacy,commercial_use=commercial_use,
                file_created_date=last_modified,source_metadata={"import_id":import_id,"client_key":key},
                relative_path=rel,manifest_id=import_id,
                supersedes_asset_id=str(existing.get("asset_id") or "") if existing and resync else "",
                project_id=project_id,
            )
            with self._lock:
                self._conn.execute(
                    "UPDATE visual_import_items SET asset_id=?,status='verified',error_code='',error_message='',updated_at=? WHERE id=?",
                    (asset["asset_id"],now(),item_id),
                )
                self._refresh_import_locked(import_id)
                self._conn.commit()
            return asset, False
        except VisualAssetError as exc:
            with self._lock:
                self._conn.execute(
                    "UPDATE visual_import_items SET status='failed',error_code=?,error_message=?,updated_at=? WHERE id=?",
                    (exc.code,str(exc)[:500],now(),item_id),
                )
                self._refresh_import_locked(import_id)
                self._conn.commit()
            raise

    def _refresh_import_locked(self, import_id: str) -> None:
        counts = self._conn.execute(
            """SELECT COUNT(*) AS seen,
                      SUM(CASE WHEN status='verified' THEN 1 ELSE 0 END) AS completed,
                      SUM(CASE WHEN status='failed' THEN 1 ELSE 0 END) AS failed
               FROM visual_import_items WHERE import_id=?""", (import_id,),
        ).fetchone()
        seen,completed,failed = (int(counts[0] or 0),int(counts[1] or 0),int(counts[2] or 0))
        total_row = self._conn.execute("SELECT total_count FROM visual_imports WHERE id=?",(import_id,)).fetchone()
        total = max(seen,int(total_row[0] or 0) if total_row else seen)
        status = "failed" if failed and not completed else "verified" if completed >= total and not failed else "pending_review"
        stamp = now()
        self._conn.execute(
            "UPDATE visual_imports SET status=?,completed_count=?,failed_count=?,updated_at=?,last_synced_at=? WHERE id=?",
            (status,completed,failed,stamp,stamp,import_id),
        )

    def get_import(self, import_id: str, user_id: str) -> dict[str, Any] | None:
        imports = self._rows("SELECT * FROM visual_imports WHERE id=? AND user_id=?",(import_id,user_id))
        if not imports:
            return None
        result = imports[0]
        result["manifest"] = loads(result.pop("manifest_json"),{})
        result["items"] = self._rows(
            "SELECT client_key,relative_path,filename,size_bytes,last_modified,asset_id,status,error_code,error_message,attempts,updated_at FROM visual_import_items WHERE import_id=? ORDER BY created_at",
            (import_id,),
        )
        return result

    def review_queue(self, user_id: str, *, status: str = "", page: int = 1, limit: int = 30) -> tuple[list[dict[str, Any]], int]:
        allowed = {"", "verified", "probable", "pending_review", "conflicted", "failed"}
        if status not in allowed:
            raise VisualAssetError("review queue 狀態無效")
        where = ["a.user_id=?"]
        params: list[Any] = [user_id]
        if status:
            where.append("a.review_status=?")
            params.append(status)
        else:
            where.append("a.review_status IN ('probable','pending_review','conflicted','failed')")
        condition = " AND ".join(where)
        with self._lock:
            total = int(self._conn.execute(f"SELECT COUNT(*) FROM visual_assets a WHERE {condition}",tuple(params)).fetchone()[0])
            rows = self._conn.execute(
                f"""SELECT a.*,COALESCE(s.canonical_name,'') AS school_name,COALESCE(c.name,'') AS club_name
                    FROM visual_assets a LEFT JOIN visual_schools s ON s.id=a.school_id
                    LEFT JOIN visual_clubs c ON c.id=a.club_id WHERE {condition}
                    ORDER BY a.updated_at DESC LIMIT ? OFFSET ?""",
                (*params,limit,(page-1)*limit),
            ).fetchall()
        return [self.public_asset(dict(row)) for row in rows],total

    def confirm_asset(self, asset_id: str, user_id: str, *, reason: str = "") -> dict[str, Any]:
        asset = self.owned_asset(asset_id,user_id)
        if not asset:
            raise VisualAssetError("找不到可確認的素材",code="not_found")
        dates = self._rows("SELECT DISTINCT value FROM visual_date_candidates WHERE asset_id=? AND status!='failed'",(asset_id,))
        if len(dates) > 1:
            with self._lock:
                self._conn.execute("UPDATE visual_assets SET review_status='conflicted',updated_at=? WHERE id=?",(now(),asset_id))
                self._conn.commit()
            raise VisualAssetError("日期證據互相衝突，請先選擇或修正日期",code="date_conflict")
        stamp = now()
        with self._lock:
            self._conn.execute(
                "UPDATE visual_assets SET review_status='verified',last_confirmed_at=?,updated_at=? WHERE id=? AND user_id=?",
                (stamp,stamp,asset_id,user_id),
            )
            self._conn.execute(
                "UPDATE visual_observations SET status='verified',updated_at=? WHERE asset_id=? AND status IN ('probable','pending_review') AND review_action!='ignored'",
                (stamp,asset_id),
            )
            self._conn.execute(
                "UPDATE visual_date_candidates SET status='verified' WHERE asset_id=? AND status='probable'",(asset_id,),
            )
            self._conn.commit()
        self.add_correction(user_id,asset_id,"asset_review",{"status":asset.get("review_status")},{"status":"verified"},reason)
        self.record_learning(user_id,"correction",asset_id=asset_id,payload={"entity_type":"asset","action":"confirm"},outcome="verified")
        return self.get_asset(asset_id,user_id) or {}

    def confirm_entity(
        self, *, user_id: str, asset_id: str, entity_type: str, action: str,
        entity_id: str = "", candidate_label: str = "", values: dict[str, Any] | None = None,
        observation_id: str = "", reason: str = "",
    ) -> dict[str, Any]:
        asset = self.owned_asset(asset_id, user_id)
        if not asset:
            raise VisualAssetError("找不到可修改的圖片", code="not_found")
        values = values or {}
        if action == "ignore":
            if not observation_id:
                raise VisualAssetError("忽略辨識結果時必須指定 observation_id")
            with self._lock:
                old = self._conn.execute("SELECT * FROM visual_observations WHERE id=? AND asset_id=?", (observation_id,asset_id)).fetchone()
                if not old:
                    raise VisualAssetError("找不到辨識結果", code="not_found")
                self._conn.execute("UPDATE visual_observations SET review_action='ignored',updated_at=? WHERE id=?", (now(),observation_id))
                self._conn.commit()
            self.add_correction(user_id,asset_id,entity_type,dict(old),{"review_action":"ignored"},reason)
            self.record_learning(user_id,"correction",asset_id=asset_id,payload={"entity_type":entity_type,"action":"ignore"},outcome="reviewed")
            return self.get_asset(asset_id,user_id) or {}
        if action not in {"confirm","correct"}:
            raise VisualAssetError("action 必須是 confirm、correct 或 ignore")

        school_id = asset.get("school_id")
        final_id = entity_id
        label = (candidate_label or values.get("name") or values.get("value") or "").strip()
        if entity_type == "person":
            if final_id:
                people = self._rows("SELECT * FROM visual_people WHERE id=? AND status='verified'", (final_id,))
                if not people:
                    raise VisualAssetError("只能比對已建立且已確認的人物", code="person_not_confirmed")
                person_row = people[0]
                if school_id and person_row.get("school_id") and school_id != person_row["school_id"]:
                    raise VisualAssetError("人物與圖片所屬學校不同，已阻止跨校混用", code="school_conflict")
                label = person_row["name"]
            else:
                if not label:
                    raise VisualAssetError("建立人物時必須提供姓名")
                person_school = self.ensure_school(str(values.get("school") or "")) or school_id
                if not person_school:
                    raise VisualAssetError("建立人物必須指定學校，避免同名人物混淆", code="school_required")
                club_id = self.ensure_club(person_school,str(values.get("club") or ""),source={"type":"user_confirmation"},confirmed=True) if values.get("club") else asset.get("club_id")
                final_id = new_id("person")
                stamp = now()
                with self._lock:
                    self._conn.execute(
                        "INSERT INTO visual_people(id,school_id,club_id,name,nickname,role,status,source,confidence,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)",
                        (final_id,person_school,club_id,label,str(values.get("nickname") or "")[:100],str(values.get("role") or "")[:120],"verified",dumps({"type":"user_confirmation","asset_id":asset_id}),1.0,stamp,stamp),
                    )
                    if not school_id:
                        self._conn.execute("UPDATE visual_assets SET school_id=?,updated_at=? WHERE id=?", (person_school,stamp,asset_id))
                    self._conn.commit()
            self.add_observation(asset_id,"person",label=label,entity_id=final_id,status="verified",confidence=1.0,source="user_confirmation",evidence={"action":action,"reason":reason})
            if values.get("use_as_reference", True):
                with self._lock:
                    self._conn.execute(
                        "INSERT OR REPLACE INTO visual_person_references(person_id,asset_id,status,created_at) VALUES(?,?,?,?)",
                        (final_id,asset_id,"verified",now()),
                    )
                    self._conn.commit()
        elif entity_type == "scene":
            if not label:
                raise VisualAssetError("請指定場景")
            final_id = self.scene_id(label)
            self.add_observation(asset_id,"scene",label=label,entity_id=final_id,status="verified",confidence=1.0,source="user_confirmation",evidence={"reason":reason})
        elif entity_type == "club":
            if not label:
                raise VisualAssetError("請指定社團")
            chosen_school = self.ensure_school(str(values.get("school") or "")) or school_id
            final_id = self.ensure_club(chosen_school,label,source={"type":"user_confirmation"},confirmed=True) or ""
            with self._lock:
                self._conn.execute("UPDATE visual_clubs SET status='verified',updated_at=? WHERE id=?", (now(),final_id))
                self._conn.execute("UPDATE visual_assets SET school_id=?,club_id=?,updated_at=? WHERE id=?", (chosen_school,final_id,now(),asset_id))
                self._conn.commit()
            self.add_observation(asset_id,"club",label=label,entity_id=final_id,status="verified",confidence=1.0,source="user_confirmation",evidence={"reason":reason})
        elif entity_type == "event":
            if not label:
                raise VisualAssetError("請指定活動名稱")
            if final_id:
                events = self._rows("SELECT * FROM visual_events WHERE id=?", (final_id,))
                if not events:
                    raise VisualAssetError("找不到活動", code="not_found")
                if school_id and events[0].get("school_id") and events[0]["school_id"] != school_id:
                    raise VisualAssetError("活動與圖片所屬學校不同，已阻止跨校混用", code="school_conflict")
                label = events[0]["name"]
            else:
                chosen_school = self.ensure_school(str(values.get("school") or "")) or school_id
                final_id = self.ensure_event(
                    school_id=chosen_school,club_id=asset.get("club_id"),name=label,
                    event_type=str(values.get("event_type") or ""),event_date=str(values.get("date") or ""),
                    start_time=str(values.get("start_time") or values.get("time") or ""),end_time=str(values.get("end_time") or ""),
                    location=str(values.get("location") or ""),status="verified",evidence=[{"source":"user_confirmation","asset_id":asset_id}],
                )
            with self._lock:
                self._conn.execute("UPDATE visual_events SET status='verified',updated_at=? WHERE id=?", (now(),final_id))
                self._conn.commit()
            self.add_observation(asset_id,"event",label=label,entity_id=final_id,status="verified",confidence=1.0,source="user_confirmation",evidence={"reason":reason})
        elif entity_type == "date":
            if not label or not normalize_date(label):
                raise VisualAssetError("日期格式無法辨識，請使用 YYYY-MM-DD")
            final_id = self.add_date(asset_id,label,source="user_input",confidence=1.0,evidence=reason,status="verified") or ""
        else:
            raise VisualAssetError("不支援的實體類型")
        stamp = now()
        with self._lock:
            self._conn.execute(
                "UPDATE visual_assets SET review_status='verified',last_confirmed_at=?,updated_at=? WHERE id=? AND user_id=?",
                (stamp,stamp,asset_id,user_id),
            )
            self._conn.commit()
        self.add_correction(user_id,asset_id,entity_type,{"candidate":candidate_label,"observation_id":observation_id},{"entity_id":final_id,"label":label,"status":"verified"},reason)
        self.record_learning(user_id,"correction",asset_id=asset_id,payload={"entity_type":entity_type,"action":action},outcome="verified")
        return self.get_asset(asset_id,user_id) or {}

    def dashboard(self, user_id: str) -> dict[str, Any]:
        visible = "(user_id=? OR privacy IN ('shared','public'))"
        with self._lock:
            total = int(self._conn.execute(f"SELECT COUNT(*) FROM visual_assets WHERE {visible} AND COALESCE(source,'')!='derived_thumbnail'",(user_id,)).fetchone()[0])
            counts = {
                "pending": int(self._conn.execute(f"SELECT COUNT(*) FROM visual_assets WHERE {visible} AND COALESCE(source,'')!='derived_thumbnail' AND review_status IN ('pending_review','probable','conflicted','failed')",(user_id,)).fetchone()[0]),
                "duplicates": int(self._conn.execute(f"SELECT COUNT(*) FROM visual_assets WHERE {visible} AND COALESCE(source,'')!='derived_thumbnail' AND duplicate_of IS NOT NULL",(user_id,)).fetchone()[0]),
                "high_quality": int(self._conn.execute(f"SELECT COUNT(*) FROM visual_assets WHERE {visible} AND COALESCE(source,'')!='derived_thumbnail' AND quality_score>=75",(user_id,)).fetchone()[0]),
                "people": int(self._conn.execute("SELECT COUNT(DISTINCT o.entity_id) FROM visual_observations o JOIN visual_assets a ON a.id=o.asset_id WHERE a.user_id=? AND o.entity_type='person' AND o.entity_id!=''",(user_id,)).fetchone()[0]),
                "clubs": int(self._conn.execute("SELECT COUNT(DISTINCT id) FROM visual_clubs WHERE id IN (SELECT club_id FROM visual_assets WHERE user_id=? AND club_id IS NOT NULL UNION SELECT o.entity_id FROM visual_observations o JOIN visual_assets a ON a.id=o.asset_id WHERE a.user_id=? AND o.entity_type='club' AND o.entity_id!='')",(user_id,user_id)).fetchone()[0]),
                "events": int(self._conn.execute("SELECT COUNT(DISTINCT o.entity_id) FROM visual_observations o JOIN visual_assets a ON a.id=o.asset_id WHERE a.user_id=? AND o.entity_type='event' AND o.entity_id!=''",(user_id,)).fetchone()[0]),
                "scenes": int(self._conn.execute("SELECT COUNT(*) FROM visual_scenes").fetchone()[0]),
            }
        recent,total_recent,_ = self.search(user_id,page=1,limit=12)
        return {"total":total,"counts":counts,"recent":recent,"recent_total":total_recent}

    def corrections(self, user_id: str, *, page: int = 1, limit: int = 30) -> tuple[list[dict[str, Any]], int]:
        with self._lock:
            total = int(self._conn.execute("SELECT COUNT(*) FROM visual_corrections WHERE user_id=?",(user_id,)).fetchone()[0])
            rows = self._conn.execute(
                "SELECT * FROM visual_corrections WHERE user_id=? ORDER BY created_at DESC LIMIT ? OFFSET ?",(user_id,limit,(page-1)*limit),
            ).fetchall()
        items = [dict(r) for r in rows]
        for item in items:
            item["old_value"] = loads(item["old_value"],{})
            item["new_value"] = loads(item["new_value"],{})
        return items,total

    def learning_insights(self, user_id: str) -> dict[str, Any]:
        rows = self._rows("SELECT event_type,COUNT(*) AS count FROM visual_learning_events WHERE user_id=? GROUP BY event_type ORDER BY count DESC",(user_id,))
        corrections = self._rows("SELECT field_name,COUNT(*) AS count FROM visual_corrections WHERE user_id=? GROUP BY field_name ORDER BY count DESC",(user_id,))
        failures = self._rows("SELECT j.error_code,COUNT(*) AS count FROM visual_analysis_jobs j JOIN visual_assets a ON a.id=j.asset_id WHERE j.status='failed' AND a.user_id=? GROUP BY j.error_code ORDER BY count DESC",(user_id,))
        return {
            "event_counts": rows, "correction_hotspots": corrections, "analysis_failures": failures,
            "policy": "只調整搜尋排序統計與策略建議；不修改模型權重，辨識事實仍需人工確認。",
        }

    def approve_correction(self, correction_id: str, approved_by: str) -> dict[str, Any] | None:
        with self._lock:
            self._conn.execute(
                "UPDATE visual_corrections SET approval_status='approved',approved_at=?,approved_by=? WHERE id=?",
                (now(),approved_by,correction_id),
            )
            self._conn.commit()
            row = self._conn.execute("SELECT * FROM visual_corrections WHERE id=?",(correction_id,)).fetchone()
        if not row:
            return None
        item = dict(row)
        item["old_value"] = loads(item["old_value"],{})
        item["new_value"] = loads(item["new_value"],{})
        return item

    def patch_asset(self, asset_id: str, user_id: str, fields: dict[str, Any]) -> dict[str, Any] | None:
        allowed = {"privacy", "commercial_use", "source", "source_metadata", "review_status", "school", "club"}
        unknown = set(fields) - allowed
        if unknown:
            raise VisualAssetError("不可修改原圖、雜湊或系統分析欄位：" + "、".join(sorted(unknown)), code="immutable_field")
        current = self.owned_asset(asset_id, user_id)
        if not current:
            return None
        updates: dict[str, Any] = {}
        if "privacy" in fields:
            if fields["privacy"] not in {"private","shared","public"}:
                raise VisualAssetError("隱私狀態無效")
            updates["privacy"] = fields["privacy"]
        if "commercial_use" in fields:
            if fields["commercial_use"] not in {"allowed","not_allowed","unknown"}:
                raise VisualAssetError("商用權限無效")
            updates["commercial_use"] = fields["commercial_use"]
        if "review_status" in fields:
            if fields["review_status"] not in STATUS_VALUES:
                raise VisualAssetError("確認狀態無效")
            updates["review_status"] = fields["review_status"]
        if "source" in fields:
            updates["source"] = str(fields["source"] or "")[:120]
        if "source_metadata" in fields:
            updates["source_metadata"] = dumps(fields["source_metadata"] or {})
        school_id = current.get("school_id")
        if "school" in fields:
            school_id = self.ensure_school(str(fields["school"] or ""))
            updates["school_id"] = school_id
            # 換學校時舊社團關聯一定清除，防止跨校污染。
            updates["club_id"] = None
        if "club" in fields:
            updates["club_id"] = self.ensure_club(school_id, str(fields["club"] or ""), source={"type":"user_correction"}, confirmed=True) if fields["club"] else None
        if updates:
            updates["updated_at"] = now()
            columns = ",".join(f"{key}=?" for key in updates)
            with self._lock:
                self._conn.execute(f"UPDATE visual_assets SET {columns} WHERE id=? AND user_id=?", (*updates.values(),asset_id,user_id))
                self._conn.commit()
        return self.get_asset(asset_id, user_id)

    def list_entities(self, kind: str, *, user_id: str, query: str = "", school: str = "", page: int = 1, limit: int = 30) -> tuple[list[dict[str, Any]], int]:
        table_map = {"people":"visual_people", "clubs":"visual_clubs", "events":"visual_events", "scenes":"visual_scenes"}
        table = table_map[kind]
        where, params = ["1=1"], []
        name_col = "name"
        if kind == "people":
            where.append("id IN (SELECT o.entity_id FROM visual_observations o JOIN visual_assets a ON a.id=o.asset_id WHERE a.user_id=? AND o.entity_type='person' AND o.entity_id!='')")
            params.append(user_id)
        elif kind == "clubs":
            where.append("id IN (SELECT club_id FROM visual_assets WHERE user_id=? AND club_id IS NOT NULL UNION SELECT o.entity_id FROM visual_observations o JOIN visual_assets a ON a.id=o.asset_id WHERE a.user_id=? AND o.entity_type='club' AND o.entity_id!='')")
            params.extend([user_id,user_id])
        elif kind == "events":
            where.append("id IN (SELECT o.entity_id FROM visual_observations o JOIN visual_assets a ON a.id=o.asset_id WHERE a.user_id=? AND o.entity_type='event' AND o.entity_id!='')")
            params.append(user_id)
        if query:
            where.append(f"{name_col} LIKE ?")
            params.append(f"%{query}%")
        if school and kind != "scenes":
            where.append("school_id IN (SELECT id FROM visual_schools WHERE canonical_name LIKE ? OR aliases LIKE ?)")
            params.extend([f"%{school}%",f"%{school}%"])
        condition = " AND ".join(where)
        with self._lock:
            total = int(self._conn.execute(f"SELECT COUNT(*) FROM {table} WHERE {condition}", tuple(params)).fetchone()[0])
            rows = self._conn.execute(
                f"SELECT * FROM {table} WHERE {condition} ORDER BY name LIMIT ? OFFSET ?",
                (*params,limit,(page-1)*limit),
            ).fetchall()
        result = [dict(r) for r in rows]
        for item in result:
            for key in ("aliases","source","social_accounts","evidence"):
                if key in item:
                    item[key] = loads(item[key], [] if key in {"aliases","evidence"} else {})
        return result,total

    def record_learning(self, user_id: str, event_type: str, *, query: str = "", asset_id: str = "", payload: dict[str, Any] | None = None, outcome: str = "") -> None:
        with self._lock:
            self._conn.execute(
                "INSERT INTO visual_learning_events(user_id,event_type,query_text,asset_id,payload,outcome,created_at) VALUES(?,?,?,?,?,?,?)",
                (user_id,event_type[:80],query[:2000],asset_id,dumps(payload or {}),outcome[:120],now()),
            )
            self._conn.commit()

    def add_to_collection(
        self, user_id: str, asset_id: str, *, kind: str, name: str = "", note: str = "", rendition: str = "original",
    ) -> dict[str, Any]:
        if kind not in {"material_pack","storyboard","social_post","poster"}:
            raise VisualAssetError("不支援的素材集合類型")
        if not self.visible_asset(asset_id,user_id):
            raise VisualAssetError("找不到圖片或沒有權限",code="not_found")
        default_names = {
            "material_pack":"目前素材包","storyboard":"目前影片分鏡",
            "social_post":"目前社群貼文","poster":"目前海報素材",
        }
        collection_name = (name or default_names[kind]).strip()[:200]
        with self._lock:
            row = self._conn.execute(
                "SELECT id FROM visual_collections WHERE user_id=? AND kind=? AND name=? AND status='active' ORDER BY updated_at DESC LIMIT 1",
                (user_id,kind,collection_name),
            ).fetchone()
            stamp = now()
            if row:
                collection_id = str(row[0])
            else:
                collection_id = new_id("collection")
                self._conn.execute(
                    "INSERT INTO visual_collections(id,user_id,name,kind,status,created_at,updated_at) VALUES(?,?,?,?,?,?,?)",
                    (collection_id,user_id,collection_name,kind,"active",stamp,stamp),
                )
            position = int(self._conn.execute(
                "SELECT COALESCE(MAX(position),-1)+1 FROM visual_collection_items WHERE collection_id=?",(collection_id,),
            ).fetchone()[0])
            self._conn.execute(
                "INSERT OR IGNORE INTO visual_collection_items(collection_id,asset_id,position,note,rendition,created_at) VALUES(?,?,?,?,?,?)",
                (collection_id,asset_id,position,note[:1000],rendition[:20],stamp),
            )
            self._conn.execute("UPDATE visual_collections SET updated_at=? WHERE id=?",(stamp,collection_id))
            self._conn.commit()
        return self.get_collection(collection_id,user_id) or {}

    def get_collection(self, collection_id: str, user_id: str) -> dict[str, Any] | None:
        rows = self._rows("SELECT * FROM visual_collections WHERE id=? AND user_id=?",(collection_id,user_id))
        if not rows:
            return None
        collection = rows[0]
        collection["items"] = self._rows(
            """SELECT i.position,i.note,i.rendition,a.id AS asset_id,a.original_filename,a.width,a.height,a.quality_score
               FROM visual_collection_items i JOIN visual_assets a ON a.id=i.asset_id
               WHERE i.collection_id=? ORDER BY i.position""",(collection_id,),
        )
        for item in collection["items"]:
            item["thumbnail_url"] = f"/api/visual-assets/{item['asset_id']}/file?variant=thumbnail"
        return collection

    def list_collections(self, user_id: str, *, kind: str = "", page: int = 1, limit: int = 30) -> tuple[list[dict[str, Any]],int]:
        where = "user_id=?" + (" AND kind=?" if kind else "")
        params: tuple[Any,...] = (user_id,kind) if kind else (user_id,)
        with self._lock:
            total = int(self._conn.execute(f"SELECT COUNT(*) FROM visual_collections WHERE {where}",params).fetchone()[0])
            rows = self._conn.execute(
                f"""SELECT c.*,(SELECT COUNT(*) FROM visual_collection_items i WHERE i.collection_id=c.id) AS item_count
                    FROM visual_collections c WHERE {where} ORDER BY updated_at DESC LIMIT ? OFFSET ?""",
                (*params,limit,(page-1)*limit),
            ).fetchall()
        return [dict(row) for row in rows],total

    def add_correction(self, user_id: str, asset_id: str, field_name: str, old_value: Any, new_value: Any, reason: str = "") -> str:
        cid = new_id("correction")
        with self._lock:
            self._conn.execute(
                "INSERT INTO visual_corrections(id,user_id,asset_id,field_name,old_value,new_value,reason,created_at) VALUES(?,?,?,?,?,?,?,?)",
                (cid,user_id,asset_id,field_name,dumps(old_value),dumps(new_value),reason[:1000],now()),
            )
            self._conn.commit()
        return cid


def infer_ratio(query: str, explicit: str = "") -> str:
    if explicit:
        return explicit
    for ratio in ("1:1","4:5","9:16","16:9"):
        if ratio in query:
            return ratio
    if "橫式" in query or "橫幅" in query:
        return "landscape"
    if "直式" in query or "限動" in query:
        return "portrait"
    if "正方形" in query:
        return "1:1"
    return ""


def infer_school(query: str) -> str:
    text = (query or "").lower()
    aliases = {
        "淡江大學": ("淡江大學","淡江","淡大","tku"),
        "國立政治大學": ("國立政治大學","政治大學","政大","nccu"),
    }
    for canonical,names in aliases.items():
        if any(name.lower() in text for name in names):
            return canonical
    return ""


def infer_club(query: str) -> str:
    """Infer only known, explicit club names from a natural-language query."""
    text = (query or "").lower()
    aliases = {
        "領袖禪學社": ("領袖禪學社", "領袖禪學", "禪學社"),
    }
    for canonical, names in aliases.items():
        if any(name.lower() in text for name in names):
            return canonical
    return ""


def ratio_matches(width: int, height: int, ratio: str) -> bool:
    value = width / max(1,height)
    if ratio == "landscape":
        return value > 1.05
    if ratio == "portrait":
        return value < 0.95
    targets = {"1:1":1.0,"4:5":0.8,"9:16":9/16,"16:9":16/9}
    target = targets.get(ratio)
    return True if target is None else abs(value-target)/target <= 0.09


def normalize_date(value: str | None) -> str:
    raw = (value or "").strip()
    if not raw:
        return ""
    raw = raw.replace("年", "-").replace("月", "-").replace("日", "").replace("/", "-").replace(".", "-")
    match = re.search(r"(?<!\d)(\d{3,4})-(\d{1,2})(?:-(\d{1,2}))?", raw)
    if not match:
        return ""
    year = int(match.group(1))
    if year < 1911:
        year += 1911
    month = int(match.group(2))
    day = int(match.group(3) or 1)
    try:
        return datetime(year, month, day).date().isoformat()
    except ValueError:
        return ""


def normalize_time(value: str | None) -> str:
    raw = (value or "").strip().replace("：",":")
    match = re.search(r"(?<!\d)(\d{1,2}):(\d{2})(?!\d)",raw)
    if not match:
        return ""
    hour,minute = int(match.group(1)),int(match.group(2))
    return f"{hour:02d}:{minute:02d}" if 0 <= hour <= 23 and 0 <= minute <= 59 else ""


_visual_store: VisualAssetStore | None = None


def get_visual_store() -> VisualAssetStore:
    global _visual_store
    desired = Path(config.DB_PATH)
    if _visual_store is None or _visual_store.path != desired:
        if _visual_store is not None:
            _visual_store.close()
        _visual_store = VisualAssetStore(desired)
    return _visual_store
