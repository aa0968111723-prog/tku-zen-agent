"""Baseline：保護既有的五個產出工具。

每個工具都要真的產生可以被對應函式庫重新開啟的檔案 ——
不接受「有回傳 ok=True 就算過」。
"""

from __future__ import annotations

from pathlib import Path

from app import tools


def test_five_core_tools_registered():
    names = {s["function"]["name"] for s in tools.SCHEMAS}
    assert {
        "search_knowledge",
        "create_spreadsheet",
        "create_document",
        "create_slides",
        "create_google_form",
    } <= names


def test_spreadsheet_produces_openable_xlsx(tmp_output_dir):
    from openpyxl import load_workbook

    r = tools.dispatch(
        "create_spreadsheet",
        {
            "filename": "測試排程",
            "sheets": [
                {"name": "社課排程", "csv": "週次,日期,主題,人數\n第1週,待填,期初茶會,35"},
                {"name": "預算", "csv": "項目,單價,數量,小計\n海報,120,10,=B2*C2"},
            ],
        },
    )
    assert r["ok"], r["message"]
    path = Path(r["local_path"])
    assert path.exists() and path.stat().st_size > 3000

    wb = load_workbook(path)
    assert wb.sheetnames == ["社課排程", "預算"]
    assert isinstance(wb["社課排程"]["D2"].value, int), "數字要存成數值型別"
    assert str(wb["預算"]["D2"].value).startswith("="), "公式要保留"
    assert wb["社課排程"].freeze_panes == "A2", "標題列要凍結"


def test_document_produces_openable_docx(tmp_output_dir):
    from docx import Document

    r = tools.dispatch(
        "create_document",
        {
            "filename": "測試文件",
            "title": "測試",
            "markdown": "# 章節\n\n**粗體**測試。\n\n- 項目一\n- 項目二\n\n| 欄A | 欄B |\n|---|---|\n| 1 | 2 |\n",
            "also_markdown": True,
        },
    )
    assert r["ok"], r["message"]
    path = Path(r["local_path"])
    doc = Document(path)
    assert len(doc.tables) == 1
    assert any("項目一" in p.text for p in doc.paragraphs)
    twins = list(path.parent.glob(f"{path.stem}-文字版*.md"))
    assert twins, "also_markdown 要另存 .md（專屬字尾，避免跨工具覆寫）"


def test_slides_produces_openable_pptx(tmp_output_dir):
    from pptx import Presentation

    r = tools.dispatch(
        "create_slides",
        {
            "filename": "測試簡報",
            "title": "封面",
            "slides_markdown": "## 第一頁\n- 重點一\n- 重點二\n\n## 第二頁\n- 另一個",
        },
    )
    assert r["ok"], r["message"]
    prs = Presentation(Path(r["local_path"]))
    assert len(prs.slides) == 3, "封面 + 兩頁"


def test_google_form_script_is_runnable_apps_script(tmp_output_dir):
    r = tools.dispatch(
        "create_google_form",
        {
            "filename": "測試表單",
            "form_title": "新生報名表",
            "questions": [
                {"type": "short_text", "title": "姓名", "required": True},
                {"type": "multiple_choice", "title": "怎麼知道我們的", "options": ["路宣", "朋友介紹"]},
                {"type": "scale", "title": "評分", "scale_min": 1, "scale_max": 5},
            ],
        },
    )
    assert r["ok"], r["message"]
    gs = Path(r["local_path"]).read_text(encoding="utf-8")
    assert "FormApp.create" in gs
    assert "addMultipleChoiceItem" in gs
    assert "addScaleItem" in gs
    assert "setDestination" in gs, "要自動建回覆試算表"
    assert "路宣" in gs, "中文選項要正確寫入"
    # extra_files 只給檔名（不洩漏伺服器絕對路徑），實際檔案在 .gs 同資料夾
    assert r["extra_files"][0].endswith(".md") and "\\" not in r["extra_files"][0]
    assert (Path(r["local_path"]).parent / r["extra_files"][0]).exists(), "要附題目一覽說明"


def test_search_knowledge_returns_grounded_context():
    r = tools.dispatch("search_knowledge", {"query": "招生管道", "top_k": 3})
    assert r["ok"]
    assert "路宣" in r["message"]
    assert "不要" in r["message"], "每次都要附不可編造的提醒"


def test_search_knowledge_flags_unknown_topic():
    r = tools.dispatch("search_knowledge", {"query": "紅燒牛肉麵的湯頭怎麼熬", "top_k": 3})
    assert r["ok"]
    assert r["low_confidence"] is True
    assert "不要自行編造" in r["message"]


# ── 容錯：開源模型常見的亂傳參數 ──────────────────────────────

def test_sheets_passed_as_bare_string_is_recovered(tmp_output_dir):
    r = tools.dispatch("create_spreadsheet", {"filename": "容錯", "sheets": "欄1,欄2\n值1,值2"})
    assert r["ok"], r["message"]


def test_sheets_passed_as_json_string_is_recovered(tmp_output_dir):
    r = tools.dispatch(
        "create_spreadsheet",
        {"filename": "容錯2", "sheets": '[{"name":"A","csv":"x,y\\n1,2"}]'},
    )
    assert r["ok"], r["message"]


def test_missing_required_arg_gives_actionable_error():
    r = tools.dispatch("create_spreadsheet", {"filename": "缺參數"})
    assert not r["ok"]
    assert "sheets" in r["message"]


def test_unknown_tool_does_not_crash():
    r = tools.dispatch("不存在的工具", {})
    assert not r["ok"]


def test_tool_exception_is_returned_not_raised(tmp_output_dir, monkeypatch):
    """工具內部炸掉時要轉成可讀訊息回給模型，不能讓整個對話中斷。"""
    from app.tools import spreadsheet

    # 簽章要跟真的一樣，dispatch 會用 inspect 檢查必要參數
    def boom(filename, sheets, destination=None, summary=""):
        raise RuntimeError("模擬爆炸")

    monkeypatch.setitem(tools._REGISTRY, "create_spreadsheet", (boom, spreadsheet.SCHEMA, "general"))
    r = tools.dispatch("create_spreadsheet", {"filename": "x", "sheets": "a,b\n1,2"})
    assert not r["ok"]
    # 新契約：例外細節只進伺服器 log，回給模型的訊息不得帶 traceback／原始例外
    assert "模擬爆炸" not in r["message"]
    assert r["message"] == "工具執行失敗，請稍後再試"
